#!/usr/bin/env python3
import json,sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
import importlib.util
HERE=Path(__file__).resolve().parent
sp=importlib.util.spec_from_file_location('g',HERE/'02_generator_prod_v02.py');g=importlib.util.module_from_spec(sp);sp.loader.exec_module(g)
c=g.load_config();rows=[]
for seed in c['calibration_seeds']:
 r=g.simulate_control(seed,c['calibration_users_control'],c);lim=c['diversity_limits']
 z=abs(r['search_zero_rate']-.4966216216)<=.03;rec=abs(r['final_recovery_rate']-.75)<=.05;h=.1390714741<=r['card_h_search_rate']<=.2655403328
 div=(r['signature']['mode_share']<=lim['signature_mode_share_max'] and r['signature']['hhi']<=lim['signature_hhi_max'] and r['signature']['shannon']>=lim['signature_shannon_min'] and r['condition_path']['mode_share']<=lim['condition_path_mode_share_max'] and r['condition_path']['hhi']<=lim['condition_path_hhi_max'] and r['condition_path']['shannon']>=lim['condition_path_shannon_min'] and r['behavior_path']['mode_share']<=lim['behavior_path_mode_share_max'] and r['behavior_path']['hhi']<=lim['behavior_path_hhi_max'] and r['behavior_path']['shannon']>=lim['behavior_path_shannon_min'] and r['template']['mode_share']<=lim['template_mode_share_max'] and r['full_path']['duplicate_rate']==0)
 r.update(pass_zero=z,pass_recovery=rec,pass_card_h=h,pass_diversity=div,simultaneous=all([z,rec,h,div]));rows.append(r)
metrics=['search_zero_rate','zero_followup_rate','immediate_recovery_rate','final_recovery_rate','card_h_search_rate','repeat_events_per_detail_search','searches_per_session']
summary={m:{'mean':sum(r[m] for r in rows)/20,'min':min(r[m] for r in rows),'max':max(r[m] for r in rows)} for m in metrics}
summary['pass_counts']={'zero':sum(r['pass_zero'] for r in rows),'recovery':sum(r['pass_recovery'] for r in rows),'card_h':sum(r['pass_card_h'] for r in rows),'diversity':sum(r['pass_diversity'] for r in rows),'simultaneous':sum(r['simultaneous'] for r in rows)}
summary['mean_pass']={'zero':abs(summary['search_zero_rate']['mean']-.4966216216)<=.01,'recovery':abs(summary['final_recovery_rate']['mean']-.75)<=.02,'card_h':abs(summary['card_h_search_rate']['mean']-.1946308725)<=.02}
wb=Workbook();ws=wb.active;ws.title='seed_results';heads=['seed']+metrics+['signature_mode','signature_hhi','signature_entropy','condition_mode','condition_hhi','condition_entropy','behavior_mode','behavior_hhi','behavior_entropy','template_mode','full_path_clones','pass_zero','pass_recovery','pass_card_h','pass_diversity','simultaneous'];ws.append(heads)
for r in rows:ws.append([r['seed']]+[r[m] for m in metrics]+[r['signature']['mode_share'],r['signature']['hhi'],r['signature']['shannon'],r['condition_path']['mode_share'],r['condition_path']['hhi'],r['condition_path']['shannon'],r['behavior_path']['mode_share'],r['behavior_path']['hhi'],r['behavior_path']['shannon'],r['template']['mode_share'],r['full_path']['duplicate_rate'],r['pass_zero'],r['pass_recovery'],r['pass_card_h'],r['pass_diversity'],r['simultaneous']])
sm=wb.create_sheet('summary');sm.append(['metric','mean','min','max']);[sm.append([m,summary[m]['mean'],summary[m]['min'],summary[m]['max']]) for m in metrics];sm.append([]);sm.append(['pass','count']);[sm.append([k,v]) for k,v in summary['pass_counts'].items()]
for sh in wb:
 for x in sh[1]:x.font=Font(bold=True,color='FFFFFF');x.fill=PatternFill('solid',fgColor='1F4E78')
 sh.freeze_panes='A2'
wb.save(HERE/'02_prod_v02_calibration_results.xlsx');print(json.dumps(summary,indent=2))
