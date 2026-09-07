#!/usr/bin/env python3
"""Read-only equivalence audit for previous and optimized 1,000-user SQLite DBs."""
import hashlib,json,math,sqlite3,sys
from collections import Counter,defaultdict
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
OLD,NEW,OUT=map(Path,sys.argv[1:4]);TABLES=['user','hotel','room','search','search_filter','search_result','event','booking']
def ro(p):c=sqlite3.connect('file:'+p.as_posix()+'?mode=ro',uri=True);c.row_factory=sqlite3.Row;c.execute('pragma query_only=on');return c
def sha(p):
 h=hashlib.sha256();
 with p.open('rb') as f:
  for b in iter(lambda:f.read(1<<20),b''):h.update(b)
 return h.hexdigest()
def norm(v):
 if not isinstance(v,str):return v
 x=' '.join(v.strip().split())
 return x[:-4] if x.endswith(' KST') else x
def wb(path,sheets):
 w=Workbook();w.remove(w.active)
 for title,heads,rows in sheets:
  s=w.create_sheet(title);s.append(heads)
  for r in rows:s.append([r.get(x) for x in heads])
  for x in s[1]:x.font=Font(bold=True,color='FFFFFF');x.fill=PatternFill('solid',fgColor='1F4E78')
  s.freeze_panes='A2'
 w.save(path)
co,cn=ro(OLD),ro(NEW);physical=[];schema=[];nulls=[];data=[];cell_diffs=[]
for label,c,p in [('previous',co,OLD),('optimized',cn,NEW)]:
 physical.append({'db':label,'path':str(p.resolve()),'bytes':p.stat().st_size,'sha256':sha(p),'sqlite_version':sqlite3.sqlite_version,'integrity':c.execute('pragma integrity_check').fetchone()[0],'quick':c.execute('pragma quick_check').fetchone()[0],'page_size':c.execute('pragma page_size').fetchone()[0],'page_count':c.execute('pragma page_count').fetchone()[0],'freelist':c.execute('pragma freelist_count').fetchone()[0],'metadata':c.execute("select count(*) from sqlite_master where type='table' and name='_generation_metadata'").fetchone()[0],'user_indexes':c.execute("select count(*) from sqlite_master where type='index' and sql is not null").fetchone()[0]})
 for t in TABLES:
  for r in c.execute(f'pragma table_info("{t}")'):schema.append({'db':label,'table':t,'column':r['name'],'type':r['type'],'notnull':r['notnull'],'default':r['dflt_value'],'pk':r['pk']})
  for col in [x['name'] for x in c.execute(f'pragma table_info("{t}")')]:
   n=c.execute(f'select count(*) from "{t}" where "{col}" is null').fetchone()[0];u=c.execute(f'select count(distinct "{col}") from "{t}"').fetchone()[0];nulls.append({'db':label,'table':t,'column':col,'null_count':n,'distinct_count':u})
for t in TABLES:
 oc=[r['name'] for r in co.execute(f'pragma table_info("{t}")')];nc=[r['name'] for r in cn.execute(f'pragma table_info("{t}")')];common=[x for x in oc if x in nc];key=common[0]
 oldrows={r[key]:tuple(r[x] for x in common) for r in co.execute(f'select {",".join(common)} from "{t}"')};newrows={r[key]:tuple(r[x] for x in common) for r in cn.execute(f'select {",".join(common)} from "{t}"')}
 keys=set(oldrows)|set(newrows);raw=sum(oldrows.get(k)!=newrows.get(k) for k in keys);normalized=sum(tuple(norm(x) for x in oldrows.get(k,()))!=tuple(norm(x) for x in newrows.get(k,())) for k in keys)
 data.append({'table':t,'old_rows':len(oldrows),'new_rows':len(newrows),'common_columns':len(common),'old_only_columns':','.join(x for x in oc if x not in nc),'new_only_columns':','.join(x for x in nc if x not in oc),'key_set_difference':len(set(oldrows)^set(newrows)),'raw_row_difference':raw,'normalized_row_difference':normalized})
 for i,col in enumerate(common):
  changed=[k for k in keys if oldrows.get(k,())[i]!=newrows.get(k,())[i]]
  normalized_changed=[k for k in keys if norm(oldrows.get(k,())[i])!=norm(newrows.get(k,())[i])]
  if changed:
   for k in sorted(changed)[:10]:cell_diffs.append({'table':t,'column':col,'changed_rows':len(changed),'normalized_changed_rows':len(normalized_changed),'business_key':k,'old_value':oldrows[k][i],'new_value':newrows[k][i]})
def metrics(c):
 total=c.execute('select count(*) from search').fetchone()[0];zero=c.execute('select count(*) from search where total_result_count=0').fetchone()[0];positive=total-zero
 sessions=c.execute('select count(distinct session_id) from search').fetchone()[0]
 recovery=c.execute("with x as(select session_id,total_result_count,row_number()over(partition by session_id order by search_time desc,search_id desc) rn,max(total_result_count=0)over(partition by session_id) z from search)select sum(rn=1 and z=1 and total_result_count>0),sum(rn=1 and z=1)from x").fetchone()
 follow=c.execute("with x as(select total_result_count,lead(search_id)over(partition by session_id order by search_time,search_id)n from search)select sum(total_result_count=0 and n is not null),sum(total_result_count=0)from x").fetchone()
 immediate=c.execute("with x as(select total_result_count,lead(total_result_count)over(partition by session_id order by search_time,search_id)n from search)select sum(total_result_count=0 and n>0),sum(total_result_count=0 and n is not null)from x").fetchone()
 hd=c.execute("select count(distinct s.search_id) from search s join search_result r on r.search_id=s.search_id and r.result_rank=1 join event e on e.search_id=s.search_id and e.hotel_id=r.hotel_id and e.event_type='hotel_detail_view' where s.total_result_count>0").fetchone()[0]
 he=c.execute("select count(*) from search s join search_result r on r.search_id=s.search_id and r.result_rank=1 join event e on e.search_id=s.search_id and e.hotel_id=r.hotel_id and e.event_type='hotel_detail_view' where s.total_result_count>0").fetchone()[0]
 return {'users':(c.execute('select count(*) from user').fetchone()[0],1),'sessions':(sessions,1),'search_zero':(zero,total),'positive_search':(positive,total),'zero_followup':follow,'immediate_recovery':immediate,'final_recovery':recovery,'card_h_search':(hd,positive),'card_h_event':(he,1),'booking':(c.execute('select count(*) from booking').fetchone()[0],1)}
mo,mn=metrics(co),metrics(cn);kpis=[]
for k in mo:kpis.append({'metric':k,'old_num':mo[k][0],'old_den':mo[k][1],'old_rate':mo[k][0]/mo[k][1],'new_num':mn[k][0],'new_den':mn[k][1],'new_rate':mn[k][0]/mn[k][1],'abs_diff':abs(mo[k][0]/mo[k][1]-mn[k][0]/mn[k][1]),'result':'PASS' if mo[k]==mn[k] else 'HOLD'})
# NULL locations on common business keys.
nulltrace=[]
for t in TABLES:
 cols=[r['name'] for r in co.execute(f'pragma table_info("{t}")') if r['name'] in {x['name'] for x in cn.execute(f'pragma table_info("{t}")')}];key=cols[0]
 for col in cols:
  a={r[0] for r in co.execute(f'select "{key}" from "{t}" where "{col}" is null')};b={r[0] for r in cn.execute(f'select "{key}" from "{t}" where "{col}" is null')}
  if a or b:nulltrace.append({'table':t,'column':col,'old_nulls':len(a),'new_nulls':len(b),'location_difference':len(a^b),'keys':','.join(map(str,sorted(a^b)[:20]))})
# Internal exact full row clones and normalized session fingerprints.
clone=[]
for t in TABLES:
 cols=[r['name'] for r in cn.execute(f'pragma table_info("{t}")')];cnt=Counter(tuple(r[x] for x in cols) for r in cn.execute(f'select * from "{t}"'));clone.append({'table':t,'exact_row_clone_excess':sum(v-1 for v in cnt.values())})
by=defaultdict(list)
for r in cn.execute('select * from search order by session_id,search_time,search_id'):by[r['session_id']].append(r)
sfmap={r['search_id']:r for r in cn.execute('select * from search_filter')};resmap=defaultdict(list);evmap=defaultdict(list)
for r in cn.execute('select * from search_result order by search_id,result_rank,hotel_id,room_id'):resmap[r['search_id']].append(r)
for r in cn.execute('select * from event where search_id is not null order by search_id,event_at,event_id'):evmap[r['search_id']].append(r)
fps=[]
for sid,ss in by.items():
 start=datetime.strptime(ss[0]['search_time'].replace(' KST',''),'%Y-%m-%d %H:%M:%S');steps=[];hotelmap={};roommap={}
 for s in ss:
  sf=sfmap[s['search_id']];res=[]
  for x in resmap[s['search_id']]:
   hotelmap.setdefault(x['hotel_id'],f'H{len(hotelmap)+1}');roommap.setdefault(x['room_id'],f'R{len(roommap)+1}');res.append((x['result_rank'],hotelmap[x['hotel_id']],roommap[x['room_id']]))
  ev=[]
  for e in evmap[s['search_id']]:
   et=datetime.strptime(e['event_at'].replace(' KST',''),'%Y-%m-%d %H:%M:%S');ev.append((e['event_type'],(et-start).total_seconds(),hotelmap.get(e['hotel_id'],'NONE')))
  steps.append(((datetime.strptime(s['search_time'].replace(' KST',''),'%Y-%m-%d %H:%M:%S')-start).total_seconds(),s['query_text'],s['checkin_date'],s['checkout_date'],s['sort_option'],s['guest_count'],s['destination'],sf['property_type'],sf['property_grade'],sf['user_rating_min'],sf['price'],sf['amenity_count'],sf['region'],s['total_result_count'],res,ev))
 fps.append(hashlib.sha256(json.dumps(steps,ensure_ascii=False,sort_keys=True,default=str).encode()).hexdigest())
fc=Counter(fps);clone3=sum(v for v in fc.values() if v>1);clone3groups=sum(v>1 for v in fc.values())
search_time_diff=next((x for x in cell_diffs if x['table']=='search' and x['column']=='search_time'),None)
classification='UNEXPLAINED_DIFFERENCE' if search_time_diff and search_time_diff['normalized_changed_rows'] else 'EXPECTED_PHYSICAL_OPTIMIZATION'
result={'physical':physical,'schema':schema,'data':data,'cell_differences':cell_diffs,'nulls':nulls,'null_trace':nulltrace,'kpis':kpis,'clone':clone,'clone3_sessions':clone3,'clone3_groups':clone3groups,'full_path_clone_excess':sum(v-1 for v in fc.values()),'schema_diff_classification':classification}
(OUT/'02_reference1000_equivalence_results_v01.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
wb(OUT/'02_reference1000_equivalence_report_v01.xlsx',[('physical',list(physical[0]),physical),('data_equivalence',list(data[0]),data),('cell_differences',list(cell_diffs[0]),cell_diffs),('kpi',list(kpis[0]),kpis),('clone',list(clone[0]),clone)])
wb(OUT/'02_reference1000_schema_diff_v01.xlsx',[('columns',list(schema[0]),schema),('objects',['db','type','name','table','sql'],[{'db':lab,'type':r[0],'name':r[1],'table':r[2],'sql':r[3]} for lab,c in [('previous',co),('optimized',cn)] for r in c.execute("select type,name,tbl_name,sql from sqlite_master order by type,name")])])
wb(OUT/'02_reference1000_text_null_audit_v01.xlsx',[('null_counts',list(nulls[0]),nulls),('null_locations',list(nulltrace[0]),nulltrace),('text_rows',list(data[0]),data),('text_differences',list(cell_diffs[0]),cell_diffs)])
co.close();cn.close();print(json.dumps({'data':data,'kpis':kpis,'null_trace':nulltrace,'clone3_sessions':clone3,'clone3_groups':clone3groups},ensure_ascii=True))
