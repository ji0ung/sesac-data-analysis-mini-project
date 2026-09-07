#!/usr/bin/env python3
"""Independent read-only clone auditor. Does not import the generator."""
import hashlib,json,math,sqlite3,sys
from collections import Counter,defaultdict
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill

QA=Path(sys.argv[1]); EXISTING=Path(sys.argv[2]); OUT=Path(sys.argv[3])
SEEDS=list(range(260601,260621)); ROUND_DIGITS=6
EXCLUDED={'user_id','session_id','search_id','event_id','search_filter_id','search_result_id','pair_id','simulation_run_id','run_id','random_seed','seed','search_at','event_at','session_start_at','session_end_at'}
TABLES=['ExperimentAssignment','Search','SearchFilter','SearchResult','ActionEvent','SearchTransition','SessionSummary']
def sha(p):
 h=hashlib.sha256()
 with p.open('rb') as f:
  for b in iter(lambda:f.read(1<<20),b''):h.update(b)
 return h.hexdigest()
def norm(v):
 if v is None:return '<NULL>'
 if isinstance(v,float):return round(v,ROUND_DIGITS)
 if isinstance(v,str):return ' '.join(v.strip().lower().split())
 return v
def dist(vals):
 f=Counter(vals);n=len(vals);ps=[x/n for x in f.values()]
 return {'mode_share':max(ps),'hhi':sum(x*x for x in ps),'entropy':-sum(x*math.log(x) for x in ps)}
def seconds(s,start):return round((datetime.fromisoformat(s)-start).total_seconds(),6)
def canonical(c,sid):
 sm=c.execute('select * from SessionSummary where session_id=?',(sid,)).fetchone();a=c.execute('select * from ExperimentAssignment where user_id=?',(sm['user_id'],)).fetchone();start=datetime.fromisoformat(sm['session_start_at']);hotel_map={};room_map={}
 def hm(x):
  if x is None:return None
  if x not in hotel_map:hotel_map[x]=f'H{len(hotel_map)+1}'
  return hotel_map[x]
 def rm(x):
  if x is None:return None
  if x not in room_map:room_map[x]=f'R{len(room_map)+1}'
  return room_map[x]
 searches=[]
 for s in c.execute('select * from Search where session_id=? order by search_seq',(sid,)):
  f=c.execute('select * from SearchFilter where search_id=?',(s['search_id'],)).fetchone();rr=[(r['result_rank'],hm(r['hotel_key']),rm(r['room_key'])) for r in c.execute('select * from SearchResult where search_id=? order by result_rank',(s['search_id'],))]
  searches.append({'seq':s['search_seq'],'t':seconds(s['search_at'],start),'result_n':s['total_result_count'],'behavior':norm(s['behavior_from_previous']),'filter':[norm(f[x]) for x in ['region_group','price_limit','option_limit_count','query_variant','continuous_preference']],'results':rr})
 events=[(norm(e['event_type']),seconds(e['event_at'],start),hm(e['hotel_key']),e['paired_mirror_event']) for e in c.execute('select * from ActionEvent where user_id=? order by event_at,event_id',(sm['user_id'],))]
 transitions=[(norm(t['behavior']),t['recovered'],t['interarrival_seconds']) for t in c.execute('select * from SearchTransition where user_id=? order by transition_id',(sm['user_id'],))]
 obj={'arm':a['sample_set_type'],'stratum':norm(a['sample_stratum']),'intent':norm(a['intent_segment']),'policy':norm(a['treatment_policy']),'searches':searches,'events':events,'transitions':transitions,'summary':[sm['search_count'],sm['experienced_zero'],sm['final_recovered'],sm['card_h_search_count'],sm['card_h_event_count'],norm(sm['outcome_segment'])]}
 raw=json.dumps(obj,ensure_ascii=False,sort_keys=True,separators=(',',':'));return hashlib.sha256(raw.encode()).hexdigest(),raw
def audit_db(p,seed):
 c=sqlite3.connect('file:'+p.as_posix()+'?mode=ro',uri=True);c.row_factory=sqlite3.Row;c.execute('pragma query_only=on')
 table=[];clone1=0;clone2g=0;clone2rows=0
 for t in TABLES:
  cols=[r['name'] for r in c.execute(f'pragma table_info("{t}")')];rows=c.execute(f'select * from "{t}"').fetchall();n=len(rows)
  full=Counter(tuple(norm(r[x]) for x in cols) for r in rows);c1=sum(v-1 for v in full.values() if v>1);clone1+=c1
  payload=[x for x in cols if x not in EXCLUDED];pc=Counter(tuple(norm(r[x]) for x in payload) for r in rows);groups=[(k,v) for k,v in pc.items() if v>1];g=len(groups);dr=sum(v for _,v in groups);clone2g+=g;clone2rows+=dr;top=max(groups,key=lambda x:x[1]) if groups else ((),0)
  table.append({'seed':seed,'table':t,'rows':n,'clone1':c1,'payload_groups':g,'payload_rows':dr,'payload_ratio':dr/n if n else 0,'top_payload':json.dumps(top[0],ensure_ascii=False),'top_n':top[1]})
 fps=defaultdict(list);structures={}
 for (sid,) in c.execute('select session_id from SessionSummary order by session_id'):
  fp,raw=canonical(c,sid);fps[fp].append(sid);structures[fp]=raw
 groups={k:v for k,v in fps.items() if len(v)>1};clone3_sessions=sum(len(v) for v in groups.values());traces=[]
 for fp,sids in groups.items():traces.append({'seed':seed,'fingerprint':fp,'sessions':','.join(sids[:2]),'session_count':len(sids),'canonical':structures[fp]})
 # KPI and existing diversity definitions.
 zero=c.execute('select avg(total_result_count=0) from Search').fetchone()[0];rec=c.execute('select avg(final_recovered) from SessionSummary where experienced_zero=1').fetchone()[0];h=c.execute("select 1.0*count(distinct e.search_id)/(select count(*) from Search where total_result_count>0) from ActionEvent e where event_type='hotel_detail_view'").fetchone()[0]
 sig=dist([r[0] for r in c.execute('select condition_signature from Search')]);conditions=[];behaviors=[];full=[];templates=[]
 for sm in c.execute('select session_id,user_id from SessionSummary order by session_id'):
  ss=list(c.execute('select condition_signature,total_result_count from Search where session_id=? order by search_seq',(sm['session_id'],)));conditions.append('>'.join(f'{x[0]}|result={int(x[1]>0)}' for x in ss));a=c.execute('select intent_segment,template_cluster_id from ExperimentAssignment where user_id=?',(sm['user_id'],)).fetchone();acts=[x[0] for x in c.execute('select behavior from SearchTransition where user_id=? order by transition_id',(sm['user_id'],))];region=c.execute('select region_group from SearchFilter f join Search s using(search_id) where s.session_id=? order by s.search_seq limit 1',(sm['session_id'],)).fetchone()[0];behaviors.append(f'{a[0][:2]}:{region[:2]}:'+'>'.join(acts));templates.append(a[1]);fp,_=canonical(c,sm['session_id']);full.append(fp)
 result={'seed':seed,'path':str(p.resolve()),'bytes':p.stat().st_size,'sha256':sha(p),'clone1':clone1,'clone2_groups':clone2g,'clone2_rows':clone2rows,'clone2_ratio':clone2rows/sum(x['rows'] for x in table),'clone3_sessions':clone3_sessions,'clone3_groups':len(groups),'full_path_clone_sessions':sum(v-1 for v in Counter(full).values()),'kpi':{'zero':zero,'recovery':rec,'card_h':h},'diversity':{'signature':sig,'condition':dist(conditions),'behavior':dist(behaviors),'template':dist(templates)}}
 c.close();return result,table,traces
# Existing XLSX is read only and is the comparison authority.
w=load_workbook(EXISTING,data_only=True,read_only=True);ws=w['seed_results'];it=ws.iter_rows(values_only=True);head=next(it);old={r[0]:dict(zip(head,r)) for r in it};w.close()
results=[];tables=[];traces=[]
for seed in SEEDS:
 r,t,x=audit_db(QA/f'calibration_control_{seed}.sqlite',seed);o=old[seed];r['kpi_match']=all(abs(r['kpi'][a]-o[b])<1e-15 for a,b in [('zero','search_zero_rate'),('recovery','final_recovery_rate'),('card_h','card_h_search_rate')]);r['diversity_match']=all(abs(r['diversity'][a][b]-o[f'{a}_{"mode" if b=="mode_share" else b}'])<1e-12 for a in ['signature','condition','behavior'] for b in ['mode_share','hhi','entropy']) and abs(r['diversity']['template']['mode_share']-o['template_mode'])<1e-12;results.append(r);tables+=t;traces+=x
def workbook(path,sheets):
 wb=Workbook();wb.remove(wb.active)
 for title,heads,rows in sheets:
  ws=wb.create_sheet(title);ws.append(heads)
  for r in rows:ws.append([r.get(x) for x in heads])
  for x in ws[1]:x.font=Font(bold=True,color='FFFFFF');x.fill=PatternFill('solid',fgColor='1F4E78')
  ws.freeze_panes='A2'
 wb.save(path)
heads=['seed','path','bytes','sha256','clone1','clone2_groups','clone2_rows','clone2_ratio','clone3_sessions','clone3_groups','full_path_clone_sessions','kpi_match','diversity_match']
workbook(OUT/'02_clone_audit_by_seed_v01.xlsx',[('by_seed',heads,results),('by_table',list(tables[0]),tables)])
workbook(OUT/'02_clone_candidate_trace_v01.xlsx',[('clone3_candidates',['seed','fingerprint','sessions','session_count','canonical'],traces if traces else [{'seed':None,'fingerprint':None,'sessions':None,'session_count':0,'canonical':'No CLONE-3 candidates'}])])
(OUT/'02_clone_audit_results_v01.json').write_text(json.dumps({'results':results,'tables':tables,'traces':traces},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'seeds':len(results),'all_kpi_match':all(x['kpi_match'] for x in results),'all_diversity_match':all(x['diversity_match'] for x in results),'clone1':sum(x['clone1'] for x in results),'clone3_sessions':sum(x['clone3_sessions'] for x in results),'full_path_clone_sessions':sum(x['full_path_clone_sessions'] for x in results)},indent=2))
