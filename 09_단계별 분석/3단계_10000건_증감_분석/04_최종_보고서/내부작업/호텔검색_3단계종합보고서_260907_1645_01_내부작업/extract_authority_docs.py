from pathlib import Path
from docx import Document
import json

ROOT = Path(__file__).resolve().parents[3]
STAGE = ROOT / "09_단계별 분석" / "3단계_10000건_증감_분석"
DOCS = [
    STAGE / "BI시각화_34일차_2팀_TO-BE실행안_교수판단회신_260904_1803_01.docx",
    ROOT / "01_프로젝트기획" / "팀프로젝트" / "2026" / "08" / "일본호텔검색_팀프로젝트계획서_20260828_v03_현행본.docx",
    ROOT / "09_단계별 분석" / "2단계_1000건_증감_분석" / "09_R3-B_독립문서검수_260905_1826_01" / "호텔검색_관측형합성1000명_전체가설분석보고서_260905_1826_03.docx",
    STAGE / "호텔검색_32일차_2팀_실행체크리스트_20260905_v04_클린제출본_[TO-BE].docx",
    STAGE / "BI시각화_33일차_2팀프로젝트실습_해답_20260905_v04_클린제출본_[TO-BE].docx",
    STAGE / "BI시각화_34일차_2팀프로젝트실습_해답_20260905_v02_클린제출본_[TO-BE].docx",
    STAGE / "BI시각화_32-34일차_2팀실습_ASIS-TOBE_세그먼트AB재설계_20260905_v02_클린제출본.docx",
    STAGE / "호텔검색_1000명_10000명_세그먼트AB시뮬레이션_증강계획서_20260905_v02_클린제출본.docx",
    STAGE / "BI시각화_32-34일차_TO-BE_슬랙공유_브리핑_20260905_v02_클린제출본.docx",
]

records = []
for path in DOCS:
    if not path.exists():
        records.append({"path": str(path), "exists": False})
        continue
    document = Document(path)
    blocks = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table_index, table in enumerate(document.tables, 1):
        for row_index, row in enumerate(table.rows, 1):
            values = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
            if any(values):
                blocks.append(f"[TABLE {table_index} ROW {row_index}] " + " | ".join(values))
    records.append({"path": str(path), "exists": True, "paragraphs_and_rows": blocks})

target = Path(__file__).with_name("authority_docs_text.json")
target.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
print(target)
