import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile
from docx import Document

ROOT=Path(__file__).resolve().parents[3]
STAGE=ROOT/'09_단계별 분석'/'3단계_10000건_증감_분석'
DOC=STAGE/'호텔검색_3단계증강및AB분석종합보고서_260907_1645_02.docx'
DB=STAGE/'호텔검색_1만명증강_탐색용생성_260907_1544_01'/'output'/'호텔검색_1만명증강_탐색용AB10000_260907_1544_01.sqlite'

def sha(path):
    h=hashlib.sha256()
    with path.open('rb') as f:
        for b in iter(lambda:f.read(1<<20),b''):h.update(b)
    return h.hexdigest()

with ZipFile(DOC) as z:
    bad=z.testzip();names=z.namelist();settings=z.read('word/settings.xml');document_xml=z.read('word/document.xml');footers=b''.join(z.read(n) for n in names if n.startswith('word/footer'))
d=Document(DOC)
texts=[p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells]
alltext='\n'.join(texts);headings=[p.text for p in d.paragraphs if p.style.name.startswith('Heading')]
required=[f'제{i}장.' for i in range(1,11)]+['부록 A.','부록 B.','부록 C.']
checks={
    'valid_zip_package':bad is None,
    'all_required_chapters':all(any(h.startswith(prefix) for h in headings) for prefix in required),
    'tables_at_least_20':len(d.tables)>=20,
    'images_at_least_6':len(d.inline_shapes)>=6,
    'update_fields_on_open':b'updateFields' in settings,
    'toc_field_present':b'TOC ' in document_xml,
    'page_field_present':b'PAGE' in footers,
    'required_limitation_statement':'실제 사용자 대상 무작위 실험 결과가 아니며' in alltext,
    'strict_hold_preserved':'strict calibration' in alltext and 'HOLD' in alltext and '16/20' in alltext,
    'card_h_history_present':'110/149=73.8%' in alltext and '29/149=19.5%' in alltext,
    'bootstrap_documented':'2,000회' in alltext and '2026090717' in alltext,
    'database_hash_present':'f6de5834ea326bb165efdd16e73ba4b1ff3939d1c8ecce27b3517d55b2a5819a' in alltext,
    'no_nan_or_none_tokens':not re.search(r'\b(?:nan|None)\b',alltext,re.I),
    'single_final_docx_in_stage':len(list(STAGE.glob('호텔검색_3단계증강및AB분석종합보고서_*.docx')))==1,
    'database_hash_unchanged':sha(DB)=='f6de5834ea326bb165efdd16e73ba4b1ff3939d1c8ecce27b3517d55b2a5819a',
}
result={
    'status':'PASS' if all(checks.values()) else 'FAIL',
    'checks':checks,
    'docx':{'path':str(DOC),'sha256':sha(DOC),'bytes':DOC.stat().st_size,'paragraphs':len(d.paragraphs),'tables':len(d.tables),'images':len(d.inline_shapes),'headings':headings,'text_characters':len(alltext),'max_table_cell_characters':max(len(c.text) for t in d.tables for row in t.rows for c in row.cells)},
    'visual_checks':{'embedded_chart_pngs_inspected':6,'chart_legibility':'PASS','full_docx_page_render':'NOT_COMPLETED','reason':'Microsoft Word COM failed even for a blank document; Computer Use native pipe unavailable. Structural Open XML and python-docx validation completed.'},
}
(Path(__file__).with_name('document_audit.json')).write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2))
