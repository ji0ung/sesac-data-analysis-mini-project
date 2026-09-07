from __future__ import annotations

import hashlib
import json
import math
import os
import platform
import sqlite3
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import yaml
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[3]
STAGE = ROOT / "09_단계별 분석" / "3단계_10000건_증감_분석"
INTERNAL = Path(__file__).resolve().parent
RUN = STAGE / "호텔검색_1만명증강_탐색용생성_260907_1544_01"
DB10 = RUN / "output" / "호텔검색_1만명증강_탐색용AB10000_260907_1544_01.sqlite"
HANDOFF = RUN / "호텔검색_1만명증강_분석용handoff_manifest_260907_1544_01.json"
QA = RUN / "호텔검색_1만명증강_전수QA결과_260907_1544_01.json"
REF = ROOT / "09_단계별 분석" / "2단계_1000건_증감_분석" / "01_초기생성_QA_전체분석" / "02_관측형합성1000명_실행묶음_260903_1606_01" / "호텔검색_관측형합성1000명_데이터_260903_1606_01_메타삭제_무인덱스_NULL5유지_텍스트최적화_16K.sqlite"
ORIGINAL = ROOT / "03_데이터모델링" / "현행데이터" / "travel_data_filtered_complete_2026-09-03_v03_비식별.sqlite"
STRICT = STAGE / "호텔검색_1만명증강_결합전이개정및검증_260907_1418_01" / "호텔검색_1만명증강_handoff_manifest_260907_1418_04.json"
CONFIG = STAGE / "호텔검색_1만명증강_결합전이개정및검증_260907_1418_01" / "호텔검색_1만명증강_config_expected_v06_260907_1418_03.yaml"
GENERATOR = STAGE / "호텔검색_1만명증강_결합전이개정및검증_260907_1418_01" / "호텔검색_1만명증강_generator_sequential_v06_260907_1418_03.py"
METRIC_CONTRACT = STAGE / "호텔검색_1만명증강_검증체인통합_260907_1326_01" / "호텔검색_1만명증강_지표정의_260907_1326_01.yaml"
AUTHORITY = STAGE / "BI시각화_34일차_2팀_TO-BE실행안_교수판단회신_260904_1803_01.docx"
OUTPUT = STAGE / "호텔검색_3단계증강및AB분석종합보고서_260907_1720_01.docx"
ANALYSIS_SEED = 2026090717
BOOTSTRAP_REPS = 2000

EXPECTED = {
    DB10: "f6de5834ea326bb165efdd16e73ba4b1ff3939d1c8ecce27b3517d55b2a5819a",
    REF: "9120561ee85705141a92eae74c5015fb2c9a20f0c8d1f6df4d99893952fd1e9f",
    ORIGINAL: "2f5bd2f73b02b103bb6107ee79aa109cf77afc6f5aecdf82b10c87c95a0bef80",
    STRICT: "32d150f92cb8c9740a4e41d9d91023e9432b5e096198dcd4a1fab2eb6284548b",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def ro(path: Path) -> sqlite3.Connection:
    connection = sqlite3.connect(f"file:{path.as_posix()}?mode=ro", uri=True)
    connection.execute("PRAGMA query_only=ON")
    return connection


def pct(value, digits=2):
    if value is None or pd.isna(value):
        return "N/A"
    return f"{100 * value:.{digits}f}%"


def pp(value, digits=2):
    if value is None or pd.isna(value):
        return "N/A"
    return f"{100 * value:+.{digits}f}%p"


def ratio(numerator, denominator):
    return numerator / denominator if denominator else np.nan


def metric_profile(path: Path) -> dict:
    connection = ro(path)
    search_table = "Search"
    event_table = "ActionEvent" if path == DB10 else "event"
    result_table = "SearchResult" if path == DB10 else "search_result"
    sequence = "search_sequence, search_time, search_id" if path == DB10 else "search_time, search_id"
    rows = connection.execute(f"""
        WITH ordered AS (
          SELECT search_id, session_id, total_result_count,
                 LEAD(search_id) OVER(PARTITION BY session_id ORDER BY {sequence}) next_search,
                 LEAD(total_result_count) OVER(PARTITION BY session_id ORDER BY {sequence}) next_count,
                 ROW_NUMBER() OVER(PARTITION BY session_id ORDER BY {sequence}) sequence_number,
                 ROW_NUMBER() OVER(PARTITION BY session_id ORDER BY {sequence.replace('search_sequence, ', '')}) chronological_number,
                 ROW_NUMBER() OVER(PARTITION BY session_id ORDER BY {sequence} DESC) reverse_number,
                 SUM(CASE WHEN total_result_count=0 THEN 1 ELSE 0 END)
                   OVER(PARTITION BY session_id ORDER BY {sequence} ROWS BETWEEN UNBOUNDED PRECEDING AND CURRENT ROW) zero_number
          FROM {search_table}
        ), details AS (
          SELECT DISTINCT e.search_id
          FROM {result_table} r JOIN {event_table} e
            ON e.search_id=r.search_id AND e.hotel_id=r.hotel_id
          WHERE r.result_rank=1 AND e.event_type='hotel_detail_view'
        )
        SELECT
          COUNT(DISTINCT o.search_id) searches,
          COUNT(DISTINCT o.session_id) sessions,
          SUM(o.total_result_count=0) zero_searches,
          SUM(o.total_result_count=0 AND o.next_search IS NOT NULL) zero_followup,
          SUM(o.total_result_count=0 AND o.next_count>0) immediate_recovery,
          SUM(o.total_result_count=0 AND o.zero_number=1 AND o.next_search IS NOT NULL) first_zero_followup,
          SUM(o.total_result_count=0 AND o.zero_number=1 AND o.next_count>0) first_zero_recovery,
          COUNT(DISTINCT CASE WHEN o.total_result_count>0 THEN o.search_id END) positive_searches,
          COUNT(DISTINCT CASE WHEN o.total_result_count>0 AND d.search_id IS NOT NULL THEN o.search_id END) card_h_searches
        FROM ordered o LEFT JOIN details d USING(search_id)
    """).fetchone()
    final = connection.execute(f"""
        WITH ordered AS (
          SELECT session_id,total_result_count,
                 ROW_NUMBER() OVER(PARTITION BY session_id ORDER BY {sequence} DESC) rn,
                 MAX(total_result_count=0) OVER(PARTITION BY session_id) ever_zero
          FROM {search_table}
        )
        SELECT SUM(rn=1 AND ever_zero=1 AND total_result_count>0),
               SUM(rn=1 AND ever_zero=1)
        FROM ordered
    """).fetchone()
    detail_events = connection.execute(f"""
        SELECT COUNT(*) FROM {search_table} s
        JOIN {result_table} r ON r.search_id=s.search_id AND r.result_rank=1
        JOIN {event_table} e ON e.search_id=s.search_id AND e.hotel_id=r.hotel_id
                            AND e.event_type='hotel_detail_view'
        WHERE s.total_result_count>0
    """).fetchone()[0]
    users = connection.execute("SELECT COUNT(*) FROM UserSynthetic" if path == DB10 else "SELECT COUNT(*) FROM user").fetchone()[0]
    active_users = connection.execute(f"SELECT COUNT(DISTINCT e.user_id) FROM {event_table} e JOIN {search_table} s USING(search_id)").fetchone()[0]
    connection.close()
    return {
        "users": users, "active_users": active_users, "sessions": rows[1], "searches": rows[0],
        "zero": (rows[2], rows[0]), "followup": (rows[3], rows[2]),
        "immediate": (rows[4], rows[3]), "first_zero_immediate": (rows[6], rows[5]),
        "final": (final[0], final[1]), "card_h": (rows[8], rows[7]),
        "card_h_events": detail_events,
    }


def session_contributions() -> pd.DataFrame:
    connection = ro(DB10)
    query = """
      WITH detail AS (
        SELECT DISTINCT e.search_id
        FROM SearchResult r JOIN ActionEvent e
          ON e.search_id=r.search_id AND e.hotel_id=r.hotel_id
        WHERE r.result_rank=1 AND e.event_type='hotel_detail_view'
      ), detail_rows AS (
        SELECT e.search_id,COUNT(*) n
        FROM SearchResult r JOIN ActionEvent e
          ON e.search_id=r.search_id AND e.hotel_id=r.hotel_id
        WHERE r.result_rank=1 AND e.event_type='hotel_detail_view'
        GROUP BY e.search_id
      ), ordered AS (
        SELECT s.*, u.pair_id, u.sample_set_type arm, u.sample_stratum,
               LEAD(s.search_id) OVER(PARTITION BY s.session_id ORDER BY s.search_sequence,s.search_time,s.search_id) next_search,
               LEAD(s.total_result_count) OVER(PARTITION BY s.session_id ORDER BY s.search_sequence,s.search_time,s.search_id) next_count,
               ROW_NUMBER() OVER(PARTITION BY s.session_id ORDER BY s.search_sequence DESC,s.search_time DESC,s.search_id DESC) rn_desc,
               SUM(CASE WHEN s.total_result_count=0 THEN 1 ELSE 0 END)
                 OVER(PARTITION BY s.session_id ORDER BY s.search_sequence,s.search_time,s.search_id ROWS BETWEEN UNBOUNDED PRECEDING AND CURRENT ROW) zero_number
        FROM Search s JOIN SessionSynthetic x USING(session_id) JOIN UserSynthetic u USING(user_id)
      )
      SELECT pair_id,arm,session_id,sample_stratum,
             COUNT(*) searches,SUM(total_result_count=0) zeros,
             SUM(total_result_count>0) positives,
             SUM(total_result_count=0 AND next_search IS NOT NULL) follow_den,
             SUM(total_result_count=0 AND next_count>0) immediate_num,
             SUM(total_result_count=0 AND zero_number=1 AND next_search IS NOT NULL) firstz_den,
             SUM(total_result_count=0 AND zero_number=1 AND next_count>0) firstz_num,
             MAX(total_result_count=0) ever_z,
             MAX(CASE WHEN rn_desc=1 AND total_result_count>0 THEN 1 ELSE 0 END) final_positive,
             SUM(total_result_count>0 AND detail.search_id IS NOT NULL) detail_searches,
             SUM(COALESCE(detail_rows.n,0)) detail_events,
             MAX(CASE WHEN search_sequence=1 AND total_result_count=0 THEN 1 ELSE 0 END) first_zero
      FROM ordered
      LEFT JOIN detail USING(search_id)
      LEFT JOIN detail_rows USING(search_id)
      GROUP BY pair_id,arm,session_id,sample_stratum
      ORDER BY pair_id,arm
    """
    frame = pd.read_sql_query(query, connection)
    connection.close()
    contribution_columns = ["searches","zeros","positives","follow_den","immediate_num","firstz_den","firstz_num","ever_z","final_positive","detail_searches","detail_events","first_zero"]
    frame[contribution_columns] = frame[contribution_columns].fillna(0)
    for item in frame["sample_stratum"].str.split("|"):
        pass
    parsed = frame["sample_stratum"].apply(lambda value: dict(part.split("=",1) for part in value.split("|")))
    frame["price_present"] = parsed.apply(lambda value: "설정" if value["price"] == "1" else "미설정")
    frame["option_bucket"] = parsed.apply(lambda value: value["opt"])
    frame["region"] = parsed.apply(lambda value: value["region"])
    frame["risk"] = parsed.apply(lambda value: value["risk"])
    frame["first_search_status"] = np.where(frame["first_zero"].eq(1), "첫 검색 0건", "첫 검색 결과 존재")
    frame["final_num"] = frame["ever_z"] * frame["final_positive"]
    frame["final_den"] = frame["ever_z"]
    frame["z_session_num"] = frame["ever_z"]
    frame["z_session_den"] = 1
    frame["session_den"] = 1
    return frame


METRICS = {
    "검색 0건률": ("zeros", "searches", "lower"),
    "0건 후속검색률": ("follow_den", "zeros", "higher"),
    "즉시 회복 전이율": ("immediate_num", "follow_den", "higher"),
    "첫 Z 즉시 회복률": ("firstz_num", "firstz_den", "higher"),
    "세션 최종 회복률": ("final_num", "final_den", "higher"),
    "Z 경험 세션률": ("z_session_num", "z_session_den", "lower"),
    "Card H-SEARCH": ("detail_searches", "positives", "higher"),
    "세션당 검색 수": ("searches", "session_den", "lower"),
}


def paired_bootstrap(frame: pd.DataFrame, reps=BOOTSTRAP_REPS, seed=ANALYSIS_SEED) -> pd.DataFrame:
    control = frame[frame.arm.eq("control")].sort_values("pair_id").reset_index(drop=True)
    treatment = frame[frame.arm.eq("treatment")].sort_values("pair_id").reset_index(drop=True)
    assert control.pair_id.tolist() == treatment.pair_id.tolist()
    rng = np.random.default_rng(seed)
    output = []
    storage = {name: [] for name in METRICS}
    n = len(control)
    for start in range(0, reps, 100):
        batch = min(100, reps-start)
        index = rng.integers(0, n, size=(batch, n), dtype=np.int32)
        for name,(num,den,_) in METRICS.items():
            cnum = control[num].to_numpy()[index].sum(axis=1)
            cden = control[den].to_numpy()[index].sum(axis=1)
            tnum = treatment[num].to_numpy()[index].sum(axis=1)
            tden = treatment[den].to_numpy()[index].sum(axis=1)
            storage[name].extend((tnum/tden - cnum/cden).tolist())
    for name,(num,den,direction) in METRICS.items():
        cn,cd = control[num].sum(),control[den].sum()
        tn,td = treatment[num].sum(),treatment[den].sum()
        cv,tv = ratio(cn,cd),ratio(tn,td)
        values=np.asarray(storage[name])
        low,high=np.quantile(values,[.025,.975])
        probability = float(np.mean(values < 0)) if direction=="lower" else float(np.mean(values > 0))
        output.append({"metric":name,"control_n":int(cn),"control_d":int(cd),"control":cv,
                       "treatment_n":int(tn),"treatment_d":int(td),"treatment":tv,
                       "difference":tv-cv,"ci_low":low,"ci_high":high,
                       "relative_change":ratio(tv-cv,cv),"ratio":ratio(tv,cv),
                       "direction":direction,"probability_favorable":probability})
    return pd.DataFrame(output)


def segment_analysis(frame: pd.DataFrame) -> pd.DataFrame:
    rng = np.random.default_rng(ANALYSIS_SEED+1)
    rows=[]
    dimensions=["price_present","option_bucket","region","risk","first_search_status"]
    labels={"price_present":"가격 조건","option_bucket":"옵션 수","region":"지역군","risk":"위험층","first_search_status":"첫 검색 상태"}
    for dimension in dimensions:
        for value in sorted(frame[dimension].unique()):
            subset=frame[frame[dimension].eq(value)]
            c=subset[subset.arm.eq("control")].sort_values("pair_id")
            t=subset[subset.arm.eq("treatment")].sort_values("pair_id")
            common=sorted(set(c.pair_id)&set(t.pair_id));c=c.set_index("pair_id").loc[common];t=t.set_index("pair_id").loc[common]
            n=len(common)
            diffs_zero=[];diffs_final=[]
            for start in range(0,BOOTSTRAP_REPS,200):
                batch=min(200,BOOTSTRAP_REPS-start);idx=rng.integers(0,n,size=(batch,n),dtype=np.int32)
                cz=c.zeros.to_numpy()[idx].sum(1)/c.searches.to_numpy()[idx].sum(1)
                tz=t.zeros.to_numpy()[idx].sum(1)/t.searches.to_numpy()[idx].sum(1)
                diffs_zero.extend((tz-cz).tolist())
                cfd=c.final_den.to_numpy()[idx].sum(1);tfd=t.final_den.to_numpy()[idx].sum(1)
                cf=np.divide(c.final_num.to_numpy()[idx].sum(1),cfd,out=np.full(batch,np.nan),where=cfd>0)
                tf=np.divide(t.final_num.to_numpy()[idx].sum(1),tfd,out=np.full(batch,np.nan),where=tfd>0)
                diffs_final.extend((tf-cf).tolist())
            dz=np.asarray(diffs_zero);df=np.asarray(diffs_final);df=df[np.isfinite(df)]
            rows.append({"dimension":labels[dimension],"segment":str(value),"pairs":n,"users":2*n,
                         "control_searches":int(c.searches.sum()),"treatment_searches":int(t.searches.sum()),
                         "control_zero":ratio(c.zeros.sum(),c.searches.sum()),"treatment_zero":ratio(t.zeros.sum(),t.searches.sum()),
                         "zero_diff":ratio(t.zeros.sum(),t.searches.sum())-ratio(c.zeros.sum(),c.searches.sum()),
                         "zero_ci_low":np.quantile(dz,.025),"zero_ci_high":np.quantile(dz,.975),
                         "control_follow_den":int(c.follow_den.sum()),"treatment_follow_den":int(t.follow_den.sum()),
                         "control_final_den":int(c.final_den.sum()),"treatment_final_den":int(t.final_den.sum()),
                         "control_final":ratio(c.final_num.sum(),c.final_den.sum()),"treatment_final":ratio(t.final_num.sum(),t.final_den.sum()),
                         "final_diff":ratio(t.final_num.sum(),t.final_den.sum())-ratio(c.final_num.sum(),c.final_den.sum()),
                         "final_ci_low":np.quantile(df,.025) if len(df) else np.nan,"final_ci_high":np.quantile(df,.975) if len(df) else np.nan,
                         "sparse":"YES" if n<50 or min(c.final_den.sum(),t.final_den.sum())<30 else "NO"})
    return pd.DataFrame(rows)


def time_profile(path: Path, arm=None) -> dict:
    connection=ro(path)
    if path==DB10:
        join="JOIN SessionSynthetic x USING(session_id) JOIN UserSynthetic u USING(user_id)"
        where="WHERE u.sample_set_type=?";params=(arm,)
        order="search_sequence,search_time,search_id"
    else:
        join="";where="";params=();order="search_time,search_id"
    values=[float(x[0]) for x in connection.execute(f"""
      WITH o AS(SELECT (julianday(search_time)-julianday(LAG(search_time) OVER(PARTITION BY session_id ORDER BY {order})))*86400.0 d FROM Search {join} {where})
      SELECT d FROM o WHERE d IS NOT NULL
    """,params) if x[0] is not None]
    connection.close();a=np.maximum(0,np.asarray(values))
    return {"n":len(a),"same":float(np.mean(a<.5)),"tail109":float(np.mean(a>109)),"p50":float(np.quantile(a,.5)),"p90":float(np.quantile(a,.9)),"p95":float(np.quantile(a,.95)),"max":float(np.max(a))}


def filter_profiles(path: Path, arm=None) -> list[dict]:
    connection=ro(path)
    if path==DB10:
        filter_table="SearchFilter"
        join="JOIN Search s USING(search_id) JOIN SessionSynthetic x USING(session_id) JOIN UserSynthetic u USING(user_id)"
        where="WHERE u.sample_set_type=?";params=(arm,)
    else:
        filter_table="search_filter"
        join="";where="";params=()
    rows=[]
    for label,expr in [("가격 설정","CASE WHEN price IS NULL THEN '미설정' ELSE '설정' END"),("옵션 수","CASE WHEN amenity_count>=2 THEN '2+' ELSE CAST(COALESCE(amenity_count,0) AS TEXT) END")]:
        q=f"SELECT {expr} category,COUNT(*) n FROM {filter_table} {join} {where} GROUP BY 1 ORDER BY 1"
        total=0;temp=[]
        for category,n in connection.execute(q,params):temp.append((category,n));total+=n
        rows.extend({"source":arm or ("reference" if path==REF else "original"),"dimension":label,"category":str(category),"n":n,"share":n/total} for category,n in temp)
    connection.close();return rows


def scenario_summary(strict: dict) -> pd.DataFrame:
    rows=[]
    for run in strict["runs"]:
        side=json.loads(Path(run["sidecar_path"]).read_text(encoding="utf-8"))
        for metric in ["search_zero_rate","immediate_recovery_transition_rate","final_recovery_rate","card_h_search_rate"]:
            cv=side["control_metrics"][metric]["value"];tv=side["treatment_metrics"][metric]["value"]
            rows.append({"scenario":run["scenario"],"seed":run["seed"],"metric":metric,"control":cv,"treatment":tv,"difference":tv-cv})
    frame=pd.DataFrame(rows)
    return frame.groupby(["scenario","metric"]).agg(runs=("seed","nunique"),control_mean=("control","mean"),treatment_mean=("treatment","mean"),diff_mean=("difference","mean"),diff_min=("difference","min"),diff_max=("difference","max")).reset_index()


def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),fill);tcPr.append(shd)


def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr();element=OxmlElement('w:tblHeader');element.set(qn('w:val'),'true');trPr.append(element)


def prevent_row_split(row):
    trPr=row._tr.get_or_add_trPr();element=OxmlElement('w:cantSplit');trPr.append(element)


def add_field(paragraph, instruction):
    run=paragraph.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=instruction;separate=OxmlElement('w:fldChar');separate.set(qn('w:fldCharType'),'separate');end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');run._r.extend([begin,instr,separate,end])


def add_table(document, headers, rows, font_size=7.5, widths=None):
    table=document.add_table(rows=1,cols=len(headers));table.style='Light Shading Accent 1';table.autofit=True
    hdr=table.rows[0];set_repeat_table_header(hdr)
    for i,h in enumerate(headers):
        hdr.cells[i].text=str(h);set_cell_shading(hdr.cells[i],'1F4E78')
        for run in hdr.cells[i].paragraphs[0].runs:run.font.color.rgb=RGBColor(255,255,255);run.font.bold=True;run.font.size=Pt(font_size)
    for row in rows:
        new_row=table.add_row();prevent_row_split(new_row);cells=new_row.cells
        for i,value in enumerate(row):
            cells[i].text="" if value is None else str(value);cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cells[i].paragraphs:
                para.paragraph_format.space_after=Pt(0)
                for run in para.runs:run.font.size=Pt(font_size);run.font.name='맑은 고딕';run._element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕')
    if widths:
        for row in table.rows:
            for i,width in enumerate(widths):row.cells[i].width=Cm(width)
    document.add_paragraph()
    return table


def caption(document,text):
    p=document.add_paragraph(text);p.style='Caption';p.alignment=WD_ALIGN_PARAGRAPH.CENTER;return p


def note(document,text):
    p=document.add_paragraph();r=p.add_run("해석 주의: ");r.bold=True;r.font.color.rgb=RGBColor(192,0,0);p.add_run(text);return p


def body(document,text,bold_prefix=None):
    p=document.add_paragraph();p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=1.12
    if bold_prefix and text.startswith(bold_prefix):p.add_run(bold_prefix).bold=True;p.add_run(text[len(bold_prefix):])
    else:p.add_run(text)
    return p


def section_guide(document, question, answer):
    """Give non-specialist readers a one-glance purpose and takeaway."""
    table=document.add_table(rows=2,cols=2);table.autofit=True
    labels=[('이 장에서 답할 질문',question,'D9EAF7'),('먼저 보는 결론',answer,'E2F0D9')]
    for row,(label,value,fill) in zip(table.rows,labels):
        prevent_row_split(row);row.cells[0].text=label;row.cells[1].text=value
        set_cell_shading(row.cells[0],fill);set_cell_shading(row.cells[1],fill)
        for run in row.cells[0].paragraphs[0].runs:
            run.bold=True;run.font.color.rgb=RGBColor(31,78,120);run.font.size=Pt(8.5)
        for run in row.cells[1].paragraphs[0].runs:run.font.size=Pt(8.5)
    table.columns[0].width=Cm(3.3);table.columns[1].width=Cm(13.0)
    document.add_paragraph()
    return table


def add_static_toc(document):
    """Visible without Word field refresh or desktop Word automation."""
    document.add_heading('목차',level=1)
    body(document,'아래 목차는 Word의 자동 필드 갱신 없이도 항상 보이는 고정 목차다. 각 장의 제목과 핵심 질문을 함께 표시해 처음 읽는 독자도 필요한 내용을 바로 찾을 수 있다.')
    rows=[
        ('먼저 읽기','프로젝트 한눈에 보기','무엇을 만들었고, 무엇을 결론 내릴 수 있는가?'),
        ('먼저 읽기','A/B 분석 핵심 결과','대조군과 실험군의 값·차이·불확실성은 얼마인가?'),
        ('제1장','핵심 요약과 사용 범위','이 보고서로 가능한 판단과 금지되는 해석은 무엇인가?'),
        ('제2장','데이터 계보와 기준 DB','실제 관측·1,000명 기준·10,000명 합성은 어떻게 다른가?'),
        ('제3장','증강 계획과 실제 생성 방법','10,000명과 A/B pair는 어떤 규칙으로 생성됐는가?'),
        ('제4장','수정·검증 과정','왜 strict HOLD인데 제한적 분석을 진행했는가?'),
        ('제5장','시나리오와 분석 질문','어떤 개선 가정을 비교했는가?'),
        ('제6장','지표·통계 방법','0건률·회복률·Card H를 어떻게 계산했는가?'),
        ('제7장','A/B 분석 결과(Control vs Treatment)','전체·세그먼트·전이별로 얼마나 달랐는가?'),
        ('제8장','인사이트와 실행 제안','현재 결과로 무엇을 우선 검증해야 하는가?'),
        ('제9장','한계와 영향','1만 명 합성 데이터가 남기는 위험은 무엇인가?'),
        ('제10장','결론과 후속 검증','실제 서비스에서 다음에 무엇을 해야 하는가?'),
        ('부록 A–C','SQL·해시·출처·미실시 분석','수치와 판단을 어떻게 재현·감사할 수 있는가?'),
    ]
    add_table(document,['구분','내용','빠르게 찾는 질문'],rows,8.2,[2.1,4.4,10.0])


def chart_setup():
    font_path=Path("C:/Windows/Fonts/malgun.ttf")
    if font_path.exists():
        from matplotlib.font_manager import FontProperties
        plt.rcParams['font.family']=FontProperties(fname=str(font_path)).get_name()
    plt.rcParams['axes.unicode_minus']=False
    plt.rcParams['figure.dpi']=160


def make_charts(effects,segments,qa,scenario,time_profiles):
    chart_setup();paths={}
    fig,ax=plt.subplots(figsize=(8.2,4.8));plot=effects.iloc[:-1].copy();y=np.arange(len(plot));x=plot.difference*100
    ax.errorbar(x,y,xerr=np.vstack([(plot.difference-plot.ci_low)*100,(plot.ci_high-plot.difference)*100]),fmt='o',color='#1f4e78',ecolor='#7f8c8d',capsize=3)
    ax.axvline(0,color='black',lw=.8);ax.set_yticks(y,plot.metric);ax.invert_yaxis();ax.set_xlabel('Treatment - Control (%p)');ax.set_title('핵심 KPI 차이와 pair bootstrap 95% 구간')
    ax.grid(axis='x',alpha=.25);fig.tight_layout();p=INTERNAL/'fig_kpi_effects.png';fig.savefig(p,bbox_inches='tight');plt.close(fig);paths['kpi']=p
    fig,ax=plt.subplots(figsize=(8.2,6.8));plot=segments.sort_values('zero_diff');labels=(plot.dimension+' · '+plot.segment).tolist();y=np.arange(len(plot));x=plot.zero_diff*100
    ax.errorbar(x,y,xerr=np.vstack([(plot.zero_diff-plot.zero_ci_low)*100,(plot.zero_ci_high-plot.zero_diff)*100]),fmt='o',color='#2e75b6',ecolor='#95a5a6',capsize=2)
    ax.axvline(0,color='black',lw=.8);ax.set_yticks(y,labels,fontsize=8);ax.set_xlabel('검색 0건률 차이 (%p, treatment - control)');ax.set_title('사전 조건 세그먼트별 탐색적 효과');ax.grid(axis='x',alpha=.25);fig.tight_layout();p=INTERNAL/'fig_segments.png';fig.savefig(p,bbox_inches='tight');plt.close(fig);paths['segment']=p
    c=qa['arm_profiles']['control']['transitions'];t=qa['arm_profiles']['treatment']['transitions'];labels=['Z→Z','Z→P','Z→END','P→Z','P→P','P→END'];keys=['Z_to_Z','Z_to_P','Z_to_END','P_to_Z','P_to_P','P_to_END']
    fig,ax=plt.subplots(figsize=(8.2,4.6));x=np.arange(6);w=.36;ax.bar(x-w/2,[c[k] for k in keys],w,label='control',color='#5b9bd5');ax.bar(x+w/2,[t[k] for k in keys],w,label='treatment',color='#ed7d31');ax.set_xticks(x,labels);ax.set_ylabel('전이/종료 건수');ax.set_title('검색 상태 전이와 종료 분포');ax.legend();ax.grid(axis='y',alpha=.2);fig.tight_layout();p=INTERNAL/'fig_transitions.png';fig.savefig(p,bbox_inches='tight');plt.close(fig);paths['transition']=p
    fig,ax=plt.subplots(figsize=(8.2,4.6));names=['기준 1,000','control','treatment'];quant=['p50','p90','p95'];x=np.arange(3);w=.24
    for i,q in enumerate(quant):ax.bar(x+(i-1)*w,[time_profiles[n][q] for n in names],w,label=q)
    ax.set_xticks(x,names);ax.set_yscale('log');ax.set_ylabel('초(log scale)');ax.set_title('검색 간격 중앙값·상위 분위수');ax.legend();ax.grid(axis='y',alpha=.2);fig.tight_layout();p=INTERNAL/'fig_time.png';fig.savefig(p,bbox_inches='tight');plt.close(fig);paths['time']=p
    mapping={'search_zero_rate':'0건률','immediate_recovery_transition_rate':'즉시 회복','final_recovery_rate':'최종 회복','card_h_search_rate':'Card H'}
    fig,axes=plt.subplots(2,2,figsize=(8.2,6));order=['conservative','expected','optimistic']
    for ax,(metric,title) in zip(axes.flat,mapping.items()):
        q=scenario[scenario.metric.eq(metric)].set_index('scenario').loc[order];x=np.arange(3);mean=q.diff_mean.values*100
        ax.errorbar(x,mean,yerr=np.vstack([(q.diff_mean-q.diff_min).values*100,(q.diff_max-q.diff_mean).values*100]),fmt='o-',capsize=3,color='#4472c4');ax.axhline(0,color='black',lw=.7);ax.set_xticks(x,['보수','기준','낙관']);ax.set_title(title);ax.set_ylabel('T-C (%p)');ax.grid(axis='y',alpha=.2)
    fig.suptitle('기존 30-seed 교정 시나리오 민감도(평균과 seed 범위)');fig.tight_layout();p=INTERNAL/'fig_scenarios.png';fig.savefig(p,bbox_inches='tight');plt.close(fig);paths['scenario']=p
    fig,ax=plt.subplots(figsize=(8.2,2.6));ax.axis('off');boxes=[('원본 v03','41 active users\n43 sessions / 296 searches','#d9eaf7'),('채택 1,000명 기준','합성 QA 기준\n1,000 sessions / 6,900 searches','#e2f0d9'),('10,000명 A/B','control 5,000 + treatment 5,000\n65,355 searches','#fff2cc')]
    for i,(title,sub,color) in enumerate(boxes):
        x=.04+i*.33;ax.add_patch(plt.Rectangle((x,.25),.26,.5,facecolor=color,edgecolor='#1f4e78'));ax.text(x+.13,.58,title,ha='center',va='center',weight='bold');ax.text(x+.13,.40,sub,ha='center',va='center',fontsize=8)
        if i<2:ax.annotate('',xy=(x+.32,.5),xytext=(x+.27,.5),arrowprops=dict(arrowstyle='->',lw=1.5))
    ax.text(.5,.08,'관측 근거 → 채택한 합성 기준 → 고정 모형의 탐색용 시나리오',ha='center',fontsize=9);fig.tight_layout();p=INTERNAL/'fig_lineage.png';fig.savefig(p,bbox_inches='tight');plt.close(fig);paths['lineage']=p
    return paths


def setup_document() -> Document:
    document=Document();section=document.sections[0];section.top_margin=Cm(1.8);section.bottom_margin=Cm(1.7);section.left_margin=Cm(1.9);section.right_margin=Cm(1.9)
    styles=document.styles
    normal=styles['Normal'];normal.font.name='맑은 고딕';normal._element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕');normal.font.size=Pt(9.5)
    normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.12
    for name,size,color in [('Title',25,'1F4E78'),('Heading 1',17,'1F4E78'),('Heading 2',13,'2F5597'),('Heading 3',11,'3F3F3F')]:
        style=styles[name];style.font.name='맑은 고딕';style._element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕');style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string(color);style.font.bold=True
    styles['Heading 1'].paragraph_format.page_break_before=True
    styles['Caption'].font.name='맑은 고딕';styles['Caption']._element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕');styles['Caption'].font.size=Pt(8);styles['Caption'].font.italic=True
    if 'Executive Quote' not in [s.name for s in styles]:
        style=styles.add_style('Executive Quote',WD_STYLE_TYPE.PARAGRAPH);style.font.name='맑은 고딕';style._element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕');style.font.size=Pt(11);style.font.color.rgb=RGBColor(31,78,120);style.font.bold=True;style.paragraph_format.left_indent=Cm(.8);style.paragraph_format.right_indent=Cm(.8);style.paragraph_format.space_before=Pt(8);style.paragraph_format.space_after=Pt(8)
    header=section.header.paragraphs[0];header.text='호텔 검색 3단계 증강 및 A/B 분석 종합보고서 · 합성 탐색 시뮬레이션';header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:run.font.size=Pt(7.5);run.font.color.rgb=RGBColor(100,100,100)
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.add_run('제한적 분석 · ');add_field(footer,'PAGE')
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');document.settings._element.append(update)
    return document


def build_document(data):
    d=setup_document();effects=data['effects'];segments=data['segments'];qa=data['qa'];scenario=data['scenario'];charts=data['charts'];config=data['config'];sources=data['sources'];ref=data['ref'];orig=data['orig'];key=effects.set_index('metric')
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(90);r=p.add_run('호텔검색 3단계 증강 및\nA/B 분석 종합보고서');r.bold=True;r.font.size=Pt(25);r.font.color.rgb=RGBColor(31,78,120)
    p=d.add_paragraph('한계 명시형 10,000명 합성 시뮬레이션 · expected 단일 run');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.runs[0].font.size=Pt(14)
    p=d.add_paragraph();p.style='Executive Quote';p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('본 보고서는 합성 기준 데이터와 가정한 개선 시나리오를 이용한 10,000명 규모의 A/B 시뮬레이션 분석이다. 실제 사용자 대상 무작위 실험 결과가 아니며, 대조군 기준선 재현성의 일부 검증 미달을 명시한 제한적 분석이다.')
    d.add_paragraph();p=d.add_paragraph(f"작성 시점: {datetime.now(ZoneInfo('Asia/Seoul')).strftime('%Y-%m-%d %H:%M KST')}");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=d.add_paragraph('상태: LIMITED_USE_READY / strict calibration: HOLD (16/20, 요구 18/20)');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=d.add_paragraph(f"분석 DB SHA-256: {EXPECTED[DB10]}");p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.runs[0].font.size=Pt(8)
    d.add_page_break();add_static_toc(d);d.add_page_break()

    p=d.add_paragraph('프로젝트 한눈에 보기');p.style='Title'
    section_guide(d,'이 프로젝트는 무엇을 해결하려는가?','호텔 검색결과 0건 이후의 회복 행동을 합성 A/B 시나리오로 탐색하되, 실제 실험 결과처럼 오해하지 않도록 근거와 한계를 함께 제시한다.')
    add_table(d,['1. 문제','2. 사용한 근거','3. 만든 데이터','4. 비교한 변화'],[[
        '검색조건이 좁으면 결과 0건이 반복되고, 어떤 완화 행동이 회복에 도움이 되는지 불명확하다.',
        '실제 비식별 원본 41 active users·43 sessions·296 searches와 별도 권위로 채택한 1,000명 합성 기준 DB.',
        '동일 사전 프로필의 5,000 pair에 control/treatment 각 1명, 총 USER·SESSION 10,000과 파생 SEARCH 65,355건.',
        'expected 가정에서 0건 후 지역 변경 중심 개입과 Z→P 회복확률 +0.06을 적용한 모형 결과.'
    ]],8.0)
    add_table(d,['핵심 결과','Treatment − Control','한 줄 해석'],[
        ['검색 0건률','−6.34%p','모형 안에서는 0건 검색이 감소했다.'],
        ['즉시 회복 전이율','+5.12%p','0건 다음 검색에서 양수 결과로 돌아오는 비율이 증가했다.'],
        ['세션 최종 회복률','+6.71%p','0건 경험 세션이 마지막에 양수 결과로 끝나는 비율이 증가했다.'],
        ['Card H-SEARCH','−0.27%p','1위 결과의 고유 검색 상세진입률은 개선 신호가 없었다.'],
    ],8.5)
    note(d,'이 네 값은 고정 seed의 합성 시뮬레이션 결과다. 실제 사용자 A/B 효과, 예약·매출 증가 또는 10,000명의 실제 관측을 뜻하지 않는다. strict 기준선 재현은 16/20으로 요구 18/20에 미달했다.')
    add_table(d,['표시','뜻','읽는 원칙'],[
        ['실제 관측','v03 비식별 원본 DB에서 직접 재계산','현실 근거지만 active 41명 소표본이다.'],
        ['합성 기준','별도 권위로 채택한 1,000명 기준 DB','실제 표본 증가로 해석하지 않는다.'],
        ['시나리오 가정','config에 넣은 개선 효과와 생성 규칙','관측값·검증된 효과로 표현하지 않는다.'],
        ['합성 결과','10,000명 DB에서 SQL로 재계산한 값','고정 모형의 민감도·방향 탐색에만 쓴다.'],
    ],8.2)
    add_table(d,['용어','쉬운 뜻'],[
        ['Z / P','검색결과 0건(Zero) / 결과 1건 이상(Positive)'],
        ['control / treatment','개선 가정을 적용하지 않은 대조군 / 적용한 실험군'],
        ['pair_id','사전 속성이 같은 control과 treatment 사용자를 연결하는 짝 ID'],
        ['Card H-SEARCH','결과 존재 검색 중 1위 호텔 상세를 한 번 이상 본 고유 검색 비율'],
        ['seed','같은 생성 결과를 다시 만들기 위한 난수 시작값'],
        ['strict HOLD','구조는 정상이어도 사전 통계 기준을 충족하지 못한 상태'],
        ['LIMITED_USE_READY','교육·발표·탐색만 허용하고 실제 성과 근거로 쓰지 않는 상태'],
    ],8.3)
    d.add_page_break()

    p=d.add_paragraph('A/B 분석 핵심 결과');p.style='Title'
    section_guide(d,'대조군과 실험군을 무엇으로 비교했고 결과는 어땠는가?','동일 사전 프로필의 5,000 pair를 비교한 결과, 합성 모형 안에서는 검색 0건이 줄고 회복이 증가했지만 상세진입은 개선되지 않았다.')
    add_table(d,['분석 설계','내용'],[
        ['비교 단위','pair_id 5,000개: 동일 사전 속성의 control 1명과 treatment 1명'],
        ['군별 규모','control 5,000명 / treatment 5,000명'],
        ['Treatment 가정','0건(Z) 후 지역 변경 중심 행동 + Z→P 회복확률 0.06 증가'],
        ['불확실성','pair 단위 bootstrap 2,000회로 조건부 95% 구간 계산'],
        ['해석 범위','고정 합성 모형의 시나리오 비교이며 실제 무작위 A/B 실험 결과가 아님'],
    ],8.4)
    ab_rows=[]
    for metric in ['검색 0건률','0건 후속검색률','즉시 회복 전이율','첫 Z 즉시 회복률','세션 최종 회복률','Z 경험 세션률','Card H-SEARCH','세션당 검색 수']:
        e=key.loc[metric]
        unit='회' if metric=='세션당 검색 수' else '%'
        if unit=='회':
            cv=f"{e.control:.4f}";tv=f"{e.treatment:.4f}";diff=f"{e.difference:+.4f}";ci=f"[{e.ci_low:+.4f}, {e.ci_high:+.4f}]"
        else:
            cv=pct(e.control,2);tv=pct(e.treatment,2);diff=pp(e.difference,2);ci=f"[{pp(e.ci_low,2)}, {pp(e.ci_high,2)}]"
        verdict=('개선 방향' if ((metric in ['검색 0건률','Z 경험 세션률','세션당 검색 수'] and e.difference<0) or (metric not in ['검색 0건률','Z 경험 세션률','세션당 검색 수','Card H-SEARCH'] and e.difference>0)) else '개선 확인 안 됨')
        if metric=='Card H-SEARCH':verdict='상세진입 개선 없음'
        ab_rows.append([metric,cv,tv,diff,ci,verdict])
    add_table(d,['A/B 지표','Control','Treatment','T−C','pair bootstrap 조건부 95% 구간','합성 모형 내 해석'],ab_rows,7.0)
    d.add_picture(str(charts['kpi']),width=Inches(6.5));caption(d,'그림 A/B-1. 핵심 지표의 Treatment−Control 차이와 조건부 95% 구간')
    add_table(d,['A/B 질문','분석 답변'],[
        ['검색 실패가 줄었는가?','예. 검색 0건률은 49.93%에서 43.59%로 6.34%p 낮았다.'],
        ['0건 이후 회복이 늘었는가?','예. 즉시 회복은 5.12%p, 세션 최종 회복은 6.71%p 높았다.'],
        ['상세 페이지 진입도 늘었는가?','아니다. Card H-SEARCH는 19.81%에서 19.54%로 0.27%p 낮아 사실상 변화가 없었다.'],
        ['실제 서비스 효과로 결론 내릴 수 있는가?','아니다. 회복 uplift가 config에 입력된 합성 시나리오이며 strict 기준선 검증도 16/20으로 HOLD다.'],
    ],8.2)
    note(d,'A/B 분석 결과는 존재하지만 “실제 A/B 실험 성과”가 아니라 “합성 조건부 비교”다. 상세 분자·분모, 기준선 재현, 세그먼트, 상태 전이, 시간 및 다중 seed 민감도는 제7장에 제시한다.')
    d.add_page_break()

    d.add_heading('제1장. 핵심 요약과 보고서의 사용 범위',level=1)
    section_guide(d,'이번 10,000명 결과를 어디까지 의사결정에 사용할 수 있는가?','데이터 구조와 가정 내 방향성은 탐색할 수 있지만, 실제 인과효과·예약·매출·일반화는 판단할 수 없다.')
    body(d,'호텔 검색의 핵심 문제는 조건이 과도하거나 탐색 범위가 좁을 때 검색결과 0건(Z)이 반복되고, 사용자가 어떤 완화 행동을 선택해야 회복되는지 불명확하다는 점이다. 1,000명 합성 기준은 이 행동 구조를 점검하는 QA 기준이며, 10,000명 확장은 실제 표본 확보가 아니라 희소 경로와 개선 시나리오의 가능한 결과를 더 안정적으로 탐색하기 위한 것이다.')
    body(d,'최종 데이터는 동일한 사전 프로필을 공유하는 pair 5,000개에 control과 treatment를 한 명씩 배정했다. 분석 대상은 USER·SESSION 각 10,000, SEARCH 65,355건이다. 통합 관계 QA는 PASS지만, expected 독립 seed의 검색 0건률 통과가 16/20으로 사전 요구 18/20에 미달해 strict calibration은 HOLD다. 별도 한계 수용 승인에 따라 교육·발표·탐색 분석만 허용된다.')
    body(d,f"고정 seed 결과에서 treatment는 control보다 검색 0건률이 {pp(key.loc['검색 0건률','difference'])} 낮고, 즉시 회복 전이율은 {pp(key.loc['즉시 회복 전이율','difference'])}, 최종 회복률은 {pp(key.loc['세션 최종 회복률','difference'])} 높았다. 반면 Card H-SEARCH는 {pp(key.loc['Card H-SEARCH','difference'])}로 실질적인 개선 신호가 없었다. 이 차이는 모형에 설정된 expected Z→P 전이 가정이 구현된 결과이지 실제 서비스 효과의 증거가 아니다.")
    add_table(d,['판단 영역','이번 보고서로 판단 가능','추가 실제 실험 필요'],[
        ['데이터 구조','관계 무결성, 시간순서, pair 균형, clone 여부','실서비스 로그 누락·추적 정확성'],
        ['시나리오 결과','고정 모형에서의 control/treatment 대비와 seed 민감도','사용자 무작위 배정에 따른 인과효과'],
        ['기능 우선순위','0건 후 회복 개입 패키지의 탐색 후보','가격·지역·옵션·검색어 개별 기여'],
        ['사업 성과','검색 회복·상세진입의 조건부 패턴','예약·매출·장기 유지 효과'],
    ],8)
    note(d,'가장 중요한 미해결점은 (1) strict 16/20, (2) 의도 세그먼트·제안 노출/선택 로그 미구현, (3) treatment가 지역 변경 중심의 복합 개입이라는 점이다. 즉시 배포보다 계측을 갖춘 소규모 실제 A/B가 다음 단계다.')

    d.add_heading('제2장. 데이터 계보와 기준 DB 채택',level=1)
    section_guide(d,'세 DB는 같은 종류의 표본인가?','아니다. 원본은 소표본 실제 로그 근거, 1,000명은 합성 QA 기준, 10,000명은 고정 모형의 시나리오 산출물이다.')
    d.add_picture(str(charts['lineage']),width=Inches(6.6));caption(d,'그림 2-1. 원본 관측 근거–합성 기준–10,000명 시뮬레이션 계보')
    add_table(d,['자료','역할','실제/합성','사용자·세션·검색','사용 목적','핵심 한계'],[
        ['v03 비식별 원본','초기 관측 근거','실제 로그 기반 비식별 DB',f"USER {orig['users']}, active {orig['active_users']} · session {orig['sessions']} · search {orig['searches']}",'문제·행동 근거','활성 41명/43세션 소표본, 로깅 편향'],
        ['채택 1,000명 기준','생성·QA 기준선','관측형 합성',f"USER {ref['users']} · session {ref['sessions']} · search {ref['searches']}",'조건부 확률·품질 기준','독립 실관측 증가 아님, 반복 패턴'],
        ['10,000명 A/B','분석 대상','expected 합성 시나리오','USER/SESSION 10,000 · search 65,355','탐색적 A/B 대비','단일 seed, 가정 의존, strict HOLD'],
        ['30-seed 교정','교정 이력','2,000명/run 합성','5+20+5 runs','재현성·민감도','최종 10,000명 DB가 아님'],
    ],7)
    body(d,'채택 기준 DB는 과거 lineage 제거 DB를 다시 포장한 동일 파일이라고 간주하지 않는다. metadata·index·물리 PK·data_origin 제거와 16K page 재패킹은 저장 구조 변경이지만, search.search_time 값 차이는 그 이력만으로 설명되지 않았다. STEP 2.8은 이 차이를 실패 사유로 덮지 않고 해당 파일을 별도 권위 기준으로 선언했다. 따라서 본 보고서는 SHA-256 912056…e9f를 독립 기준으로 사용하며 과거 DB와 시간값 동등성을 주장하지 않는다.')
    body(d,'기준 DB의 물리 PK·인덱스 부재는 논리키가 필요 없다는 뜻이 아니다. 고유키 NULL·중복, 참조 orphan, 검색–필터–결과–이벤트 연결을 SQL로 별도 검증했다. 10,000명 DB는 별도 schema에서 물리 키와 관계를 구성하므로 두 DB의 저장 구조를 혼동하지 않는다.')
    add_table(d,['전부 NULL 필드','기준 DB 위치','10,000명 대응 필드','분석 제약'],[
        ['event.review_completed_at','event','ActionEvent.review_completed_at','리뷰 완료 시점 분석 불가'],
        ['event.review_text','event','ActionEvent.review_text','리뷰 내용·감성 분석 불가'],
        ['search_filter.property_type','search_filter','SearchFilter.property_type','숙소 유형별 효과 분석 불가'],
        ['search_filter.property_grade','search_filter','SearchFilter.property_grade','등급 조건별 효과 분석 불가'],
        ['user.age_group','user','UserSynthetic.age_group','연령별 이질성 분석 불가'],
    ],8)
    body(d,'기준과 생성 데이터의 날짜·시간은 시뮬레이션 순서와 간격을 표현하는 값이다. 실제 서비스 발생 시점, 계절성, 요일 효과로 해석할 수 없다.')

    d.add_heading('제3장. 증강 계획과 실제 생성 방법',level=1)
    section_guide(d,'10,000명은 단순 복제인가, 새 실제 관측인가?','둘 다 아니다. pair 설계로 새 합성 경로를 생성했지만 현실 표본이 늘어난 것은 아니며, 결과는 config 가정에 의존한다.')
    d.add_heading('3.1 증강 단위와 paired 설계',level=2)
    body(d,'최종 10,000명은 기준 1,000명에 9,000명을 덧붙인 11,000명 표본이 아니다. 새 UserSynthetic 10,000명과 SessionSynthetic 10,000개를 만들고, 검색·결과·이벤트 행 수는 각 상태 경로에서 파생했다. pair_id 하나에 control/treatment 각 1명이 대응하며 sample_stratum·source_session_id 등 사전 속성은 pair 내 동일하다.')
    body(d,'공통 random stream은 지역군·가격 설정·옵션 수·위험층·첫 검색 조건을 공유하고, arm별 state stream은 그 뒤의 경로를 분리한다. 사전 pair 속성 불일치와 treatment leakage는 모두 0건이다. source_session_id는 기준 세션을 직접 복제한 템플릿 ID가 아니라 무작위로 연결한 계보 표식이며, 검색 상태와 결과 호텔은 별도로 추출한다.')
    d.add_heading('3.2 대조군·실험군 상태 전이',level=2)
    body(d,'대조군과 실험군 모두 첫 검색을 P(결과 존재) 또는 Z(0건)로 시작하고, 현재 상태·첫 Z 여부·이전 Z 경험에 따라 Z/P/END로 이동한다. 세션은 END 추출 또는 최대 31검색에서 종료된다. positive result_count와 호텔 목록은 기준 DB의 양수 결과 수 분포 및 1,000개 공통 호텔 차원에서 추출한다.')
    body(d,'expected treatment의 실제 변경은 Z 상태에서 다음 P로 회복할 확률에 +0.06을 더하고, 직전 검색이 Z이면 후속 action_type을 region_change로 기록하는 것이다. 계획 문서의 가격 범위 확대·옵션 감소·동일 조건 반복 방지·검색어 제안은 개별 제안 노출/선택이나 6개 intent_segment로 분리 구현되지 않았다. 그러므로 본 분석은 패키지 시나리오만 비교하며 개별 기능의 기여를 식별하지 않는다.')
    d.add_heading('3.3 결과·상세·시간 생성',level=2)
    body(d,f"결과 존재 검색은 Card H-SEARCH 확률 {config['card_h']['search_probability']:.6f}로 1위 호텔 상세진입 여부를 생성하고, 선택된 검색에는 평균 반복량 설정 {config['card_h']['repeat_mean']:.4f}를 이용해 click/detail 이벤트 쌍을 만든다. Card H는 고유 검색 기준이며 반복 이벤트량을 전환율 분자에 중복 합산하지 않는다.")
    tm=config['time_mixture'];body(d,f"검색 간격은 동일 timestamp 확률 {tm['same_timestamp_probability']:.4%}, body 로그정규, tail 확률 {tm['tail_probability']:.1%}의 혼합이다. body는 {tm['body_cap_seconds']}초, 전체는 {tm['maximum_seconds']:,}초로 cap한다. 각 검색 후 impression·click/detail·session_end가 시간순으로 배치되며 역전은 0건이다. 긴 꼬리를 보존했지만 실제 체류시간이나 네트워크 지연을 측정한 것은 아니다.")
    d.add_heading('3.4 복제 방지와 재현성',level=2)
    body(d,'CLONE-A는 동일 행, CLONE-B는 ID·절대시간을 제외한 동일 세션, CLONE-C는 결과 호텔·순위까지 포함한 완전 경로 복제를 검사한다. DIVERSITY-PATH는 추상 상태 경로 집중도를 설명하며 clone 판정과 다르다. 짧은 경로와 delay-vector는 가능한 조합이 제한되어 자연 충돌할 수 있다.')
    body(d,'고정 seed 2434815518에서 내부 business signature 충돌은 treatment 4회 발생해 승인된 방식으로 재추첨되었고 실패는 0회였다. 재추첨은 완전 경로 중복을 막지만 발생 확률이 매우 낮아 전체 분포 영향은 제한적이다. 결정론적 replay에서 검색 수와 0건 수가 DB와 완전히 일치했다.')
    add_table(d,['설계 항목','원래 계획','실제 구현','검증 근거','남은 한계'],[
        ['표본','control/treatment 5,000씩','5,000씩, pair 5,000','전수 QA','실관측 아님'],
        ['의도','6개 intent_segment','별도 필드 없음; sample_stratum 사용','schema·DB 프로파일','의도별 기능 효과 불가'],
        ['처치','가격·옵션·지역·반복·검색어 맞춤','Z→P +0.06, Z 후 region_change 중심','v06 코드·config','패키지 효과만 가능'],
        ['노출 퍼널','노출→선택→변경','session treatment_exposed만 존재','필드 검사','선택·수락률 불가'],
        ['시간','body/tail 혼합','동일 timestamp+로그정규+cap','시간 QA','실제 체류시간 아님'],
        ['복제 방지','행·세션·경로 복제 금지','CLONE-A/B/C=0, 4회 재추첨','clone audit/replay','현실 대표성 보장 아님'],
        ['예약','검증 후 별도','BOOKING 0행','row count','매출 효과 불가'],
        ['용량','과거 150MB 목표','790.736896MB full-detail','파일 실측','저장비용 큼'],
    ],7)

    d.add_heading('제4장. 수정·검증 과정과 한계 명시형 사용 전환',level=1)
    section_guide(d,'왜 검증을 거쳤는데도 HOLD인가?','관계·시간·clone QA는 통과했지만, 독립 seed 기준선 재현은 16/20으로 사전 요구 18/20에 미달했기 때문이다.')
    add_table(d,['단계','발견한 문제','수정 내용','검증 방법','결과','남은 문제'],[
        ['Card H 정정','이벤트 110행/검색 149건 grain 혼용','고유 검색 29/149를 H-SEARCH, 110행을 H-EVENT로 분리','원본 SQL trace','정의 종결','과거 73.83%와 비교 불가'],
        ['회복 정의','558/3271과 69/628 혼용','모든 Z 전이와 세션 첫 Z 전이를 별도 지표화','공통 contract SQL','분리 완료','사후 분모 주의'],
        ['상태 전이','길이·회복·종료를 독립 설정해 drift','history-aware Z/P/END 결합','seed별 재계산','대부분 재현','0건률 16/20'],
        ['clone 정의','추상 경로 반복을 실제 복제로 오판','CLONE-A/B/C와 DIVERSITY-PATH 분리','approved clone audit','실제 clone 0','짧은 경로 자연 반복'],
        ['검증 연결','metric/clone gate 연결 누락','공통 calculator·adapter·통합 gate 연결','30개 DB gate','30/30 구조 PASS','통계 재현성은 별도'],
        ['승인 인터페이스','calibration authorization 전달 문제','동결 wrapper·manifest 경로 정정','해시·회귀 테스트','인터페이스 PASS','production 승인은 아님'],
        ['독립 평가','expected 18/20 필요','20 seed 동결 평가','분자/분모·hash 검산','16/20 HOLD','기준선 seed 민감도'],
        ['제한 사용','정식 PASS 불가','HOLD 보존+교육/탐색 예외 승인','10k 전수 QA','LIMITED_USE_READY','실제 A/B 필요'],
    ],6.7)
    body(d,'교정 expected 20개 seed의 control 검색 0건률은 16/20이 사전 구간을 통과했고 요구는 18/20이었다. 기준 DB는 3,434/6,900=49.7681%, pooled control은 69,173/137,731=50.2233%로 +0.4551%p 높았으며 seed 범위는 48.7615%~51.5865%였다. 16/20은 seed별 판정 통과 수이지 데이터 20%가 오류라는 뜻이 아니다. 표본 규모가 커지면 같은 모형 안의 Monte Carlo 오차는 줄 수 있지만 모형 편향은 자동으로 사라지지 않는다.')
    body(d,'최종 10,000명 control의 0건률은 16,761/33,572=49.9255%로 기준 대비 +0.1574%p이며 기존 seed 관측 범위 안이다. 이 한 번의 결과로 16/20 HOLD를 소급 해제하지 않는다.')

    d.add_heading('제5장. 시나리오와 분석 질문',level=1)
    section_guide(d,'무엇을 바꿨고 어떤 질문을 검증했는가?','expected 단일 대표 run을 중심으로 0건·회복·상세진입의 방향을 비교하고, 과거 30-seed 결과는 민감도 근거로만 사용한다.')
    add_table(d,['개선안','근거 문서·카드','대상 조건','계획상 개입','실제 구현','검증 지표','근거 한계'],[
        ['가격 범위 확대','교수 우선순위 1·카드 A','가격 설정','가격 상향/해제','개별 노출 없음','가격 stratum 조건부 KPI','개별 기여 식별 불가'],
        ['옵션 수 축소','우선순위 2·카드 A','옵션 2+','요구 수 감소','개별 노출 없음','옵션 stratum 조건부 KPI','종류 필드 없음'],
        ['인접 지역 확대','우선순위 3·카드 D','위치 유연','region_change','Z 후 treatment에 강제','Z→P·최종 회복','복합 uplift와 결합'],
        ['동일 조건 반복 방지','우선순위 4·카드 F','조건 고수','반복 경고','전용 플래그 없음','same_condition 행동','의도군 없음'],
        ['단계형 완화','우선순위 5, 근거 확인 필요','빠른 해결','단계별 완화','별도 정책 로그 없음','세션 길이·회복','개별 단계 식별 불가'],
        ['연관어 제안','우선순위 6·검색어 n=10','표현 수정','query_change','전용 노출 없음','query 행동·Card H','근거 매우 희소'],
    ],6.5)
    body(d,'conservative/expected/optimistic의 실제 v06 처치 강도는 Z→P 확률에 각각 +0.03/+0.06/+0.09를 더하는 가정이다. 관측 추정값이나 95% 구간이 아니며, 기존 treatment_posteriors와 search_hazard 항목은 최종 sequential_plan의 처치 전이에 직접 사용되지 않는다. 본 10,000명 DB는 expected 한 run이며 새 시나리오 DB는 생성하지 않았다.')
    add_table(d,['기존 질문','현재 분석 질문','지표','분석 단위','가능 여부','해석 범위'],[
        ['A1/A2 제한조건','가격/옵션 사전조건별 차이가 다른가','0건률·최종회복','search/session/pair','조건부 가능','의도·개별 기능 효과 아님'],
        ['B1 후속검색','Z 이후 검색이 이어지는가','후속검색률','Z search','가능','이탈 직접 관측 아님'],
        ['B2/B3 회복','패키지가 즉시/최종 회복을 바꾸는가','즉시·최종 회복률','transition/session','가능','합성 가정 조건부'],
        ['H3 상세진입','회복/결과 존재 뒤 상세가 다른가','Card H-SEARCH','positive search','가능','전체 Card H, 회복자만 효과 아님'],
        ['TO-BE H0/H1','전체 T-C 차이','핵심 KPI','paired assignment','가능','실제 인과효과 아님'],
        ['TO-BE H2','6개 의도군 이질성','intent×arm','user/pair','불가','intent_segment 없음'],
        ['TO-BE H4~H6','개별 기능 효과','노출·선택·행동','exposure/user','불가','제안 로그·전용 treatment 없음'],
        ['예약 성과','예약/매출 변화','booking rate/revenue','user/session','불가','BOOKING 0행'],
    ],7)

    d.add_heading('제6장. 지표 정의와 분석 방법',level=1)
    section_guide(d,'서로 다른 분모와 grain을 어떻게 구분했는가?','검색·전이·세션·이벤트 grain을 분리하고, Card H는 반복 이벤트가 아닌 고유 search_id 기준으로 계산했다.')
    add_table(d,['지표 ID','질문','단위','분자','분모','포함·제외','계산·주의'],[
        ['M1 검색 0건률','검색 실패 비중','search_id','result_count=0 고유 검색','전체 고유 검색','두 arm 별도','행 수가 아닌 고유 검색'],
        ['M2 후속검색률','Z 뒤 검색 지속','Z search','다음 검색 존재 Z','전체 Z','세션 마지막 Z 포함 분모','이탈 직접 관측 아님'],
        ['M3 즉시 회복','Z 다음 P','Z transition','다음 결과>0 Z','후속 있는 Z','모든 Z','첫 Z 지표와 분리'],
        ['M4 첫 Z 회복','세션 첫 Z 다음 P','session first-Z','첫 Z 다음 결과>0','후속 있는 세션 첫 Z','세션당 1건','M3와 grain 다름'],
        ['M5 최종 회복','마지막이 P인가','session_id','마지막 결과>0 Z경험 세션','Z 경험 세션','never-Z 제외','한번이라도 회복과 다름'],
        ['M6 Z 경험','실패 경험 세션','session_id','Z≥1 세션','전체 세션','arm별','사후 상태'],
        ['M7 Card H-SEARCH','1위 상세 진입','positive search_id','rank1 detail≥1 검색','결과 존재 검색','반복 1회 처리','click 별도 합산 금지'],
        ['M8 Card H-EVENT','반복 탐색 강도','event row','rank1 detail 행','해당 없음','보조량','비율 아님'],
    ],6.8)
    add_table(d,['채택 1,000명 기준 지표','분자/분모','비율'],[
        ['검색 0건률','3,434/6,900','49.7681%'],['0건 후 후속검색률','3,271/3,434','95.2533%'],['즉시 회복 전이율','558/3,271','17.0590%'],['세션별 첫 Z 즉시 회복률','69/628','10.9873%'],['최종 회복률','488/651','74.9616%'],['Card H-SEARCH','680/3,466','19.6192%'],['Card H-EVENT','2,580행','반복량'],
    ],8)
    body(d,'Card H 변경 이력: 교수 회신은 당시 승인값 110/149=73.8%와 새 29/149=19.5%의 불일치를 종결하라고 요구했다. 이후 원본 SQL 재검산에서 110은 반복 event row, 29는 detail이 1회 이상 존재하는 고유 search임이 확인됐다. 현행 contract는 H-SEARCH=29/149와 H-EVENT=110행을 분리하고 73.83%를 전환율에서 폐기했다. 채택 1,000명 기준은 같은 고유검색 정의로 680/3,466이다.')
    body(d,f'불확실성은 pair_id 5,000개를 단위로 {BOOTSTRAP_REPS:,}회 복원추출했다. 한 pair가 선택되면 두 arm의 모든 검색·이벤트 기여를 함께 합산하고 매 반복에서 분자/분모를 다시 계산했다. 분석 seed는 {ANALYSIS_SEED}다. 제시 구간은 고정 합성 모형과 이번 생성 결과에 조건부인 bootstrap 구간이며 원본 편향·모형 오지정·현실 treatment 효과 불확실성을 포함하지 않는다.')
    body(d,'세그먼트는 개입 전 sample_stratum의 가격 설정, 옵션 수, 지역군, risk 및 pair 공통 첫 검색 상태만 사용했다. SG1~SG4와 회복 여부는 사후 결과이므로 사전 세그먼트로 사용하지 않았다. 여러 세그먼트 비교는 탐색적이며 다중비교 보정을 통한 확증 판단을 하지 않는다.')

    d.add_heading('제7장. A/B 분석 결과 — Control vs Treatment',level=1)
    section_guide(d,'합성 control과 treatment 사이에 어떤 차이가 나타났는가?','0건률은 낮고 회복률은 높았지만 이는 설정한 전이 가정의 결과다. Card H는 사실상 변하지 않았다.')
    d.add_heading('7.1 규모와 전수 품질',level=2)
    add_table(d,['검사','결과','판정'],[
        ['USER / SESSION','10,000 / 10,000','PASS'],['control / treatment / pair','5,000 / 5,000 / 5,000','PASS'],['SEARCH / RESULT / EVENT','65,355 / 2,023,419 / 2,152,148','정보'],['integrity_check / FK','ok / 0','PASS'],['시간 역전 / 종료 후 이벤트','0 / 0','PASS'],['결과 수·노출 연결 위반','0','PASS'],['pair 사전속성 불일치 / leakage','0 / 0','PASS'],['CLONE-A/B/C / reference session clone','0 / 0 / 0 / 0','PASS'],['BOOKING / NULL5 위반','0 / 0','PASS'],['raw gate / limited status','PASS / LIMITED_USE_READY','제한 승인'],['strict calibration','16/20 (요구 18/20)','HOLD'],
    ],8)
    d.add_heading('7.2 기준선 재현',level=2)
    rows=[]
    mapping=[('검색 0건률','zero'),('0건 후속검색률','followup'),('즉시 회복 전이율','immediate'),('첫 Z 즉시 회복률','first_zero_immediate'),('세션 최종 회복률','final'),('Card H-SEARCH','card_h')]
    for label,keyname in mapping:
        rr=data['ref'][keyname];e=effects[effects.metric.eq(label)].iloc[0]
        rows.append([label,f"{rr[0]:,}/{rr[1]:,}",pct(ratio(*rr),4),f"{e.control_n:,}/{e.control_d:,}",pct(e.control,4),pp(e.control-ratio(*rr),4)])
    add_table(d,['지표','기준 n/d','기준','10k control n/d','control','차이'],rows,7.5)
    body(d,'control 검색 0건률의 기준 편차는 +0.1574%p로 교정 pooled 편차 +0.4551%p보다 작지만, 이것은 seed 한 번의 실현값이다. 세션당 검색은 6.7144로 기준 6.9000보다 0.1856회(-2.69%) 적어 검색 grain의 노출량이 감소했다. 즉 핵심 비율은 대체로 기준 근처이나 검색량 구조가 완전히 동일하지는 않다.')
    d.add_heading('7.3 전체 배정군 A/B 비교',level=2)
    rows=[]
    for _,e in effects.iterrows():
        rows.append([e.metric,f"{e.control_n:,}/{e.control_d:,}",pct(e.control,4),f"{e.treatment_n:,}/{e.treatment_d:,}",pct(e.treatment,4),pp(e.difference,4),f"[{pp(e.ci_low,3)}, {pp(e.ci_high,3)}]",f"{e.ratio:.3f}",pct(e.probability_favorable,1)])
    add_table(d,['지표','C n/d','C','T n/d','T','T−C','조건부 95% 구간','비율비','유리 확률'],rows,6.5)
    d.add_picture(str(charts['kpi']),width=Inches(6.6));caption(d,'그림 7-1. 전체 배정군 핵심 KPI 차이와 pair bootstrap 구간')
    body(d,f"검색 0건률은 control {pct(key.loc['검색 0건률','control'],4)}에서 treatment {pct(key.loc['검색 0건률','treatment'],4)}로 {pp(key.loc['검색 0건률','difference'],4)} 변했다. 즉시 회복은 {pp(key.loc['즉시 회복 전이율','difference'],4)}, 최종 회복은 {pp(key.loc['세션 최종 회복률','difference'],4)}다. 모형 안에서는 회복 방향이 일관되지만, treatment_recovery_uplift를 직접 설정했으므로 이 결과는 가정의 작동 확인이다.")
    body(d,f"Card H-SEARCH는 control {pct(key.loc['Card H-SEARCH','control'],4)}, treatment {pct(key.loc['Card H-SEARCH','treatment'],4)}로 {pp(key.loc['Card H-SEARCH','difference'],4)}였다. Card H 생성확률은 arm 공통이어서 상세진입 개선을 기대하도록 설계되지 않았으며, 실제 차이도 작다. 따라서 회복 증가가 상세진입 증가로 이어졌다고 결론 내릴 수 없다.")
    note(d,'사후 Z 경험 세션끼리의 최종 회복률은 조건부 기술통계다. treatment가 Z 경험 집단 구성 자체를 바꿀 수 있으므로 이를 전체 사용자 평균 처치효과와 동일시하지 않는다.')
    d.add_heading('7.4 사전 조건별 탐색적 차이',level=2)
    segrows=[]
    for _,s in segments.iterrows():segrows.append([s.dimension,s.segment,f"{s.pairs:,}",f"{s.control_searches:,}/{s.treatment_searches:,}",pp(s.zero_diff),f"[{pp(s.zero_ci_low)}, {pp(s.zero_ci_high)}]",f"{s.control_final_den:,}/{s.treatment_final_den:,}",pp(s.final_diff),s["sparse"]])
    add_table(d,['구분','세그먼트','pair','검색 C/T','0건률 차이','조건부 95% 구간','최종회복 분모 C/T','최종회복 차이','희소'],segrows,6.2)
    d.add_picture(str(charts['segment']),width=Inches(6.5));caption(d,'그림 7-2. 사전 조건 세그먼트별 검색 0건률 차이')
    body(d,'대부분 사전 조건에서 treatment의 0건률이 낮아지는 방향이지만 효과 크기는 지역·옵션·첫 검색 상태에 따라 다르다. 이는 실제 기능별 이질성이 아니라 동일한 Z→P uplift가 서로 다른 경로 길이·Z 노출 빈도에 작용한 조건부 결과일 수 있다. intent_segment가 없으므로 계획된 6개 의도군의 효과 순위를 만들지 않았다.')
    d.add_heading('7.5 상태 전이와 종료',level=2)
    c=qa['arm_profiles']['control']['transitions'];t=qa['arm_profiles']['treatment']['transitions']
    add_table(d,['전이','control','treatment','차이'],[[label,f"{c[k]:,}",f"{t[k]:,}",f"{t[k]-c[k]:+,}"] for label,k in zip(['Z→Z','Z→P','Z→END','P→Z','P→P','P→END'],['Z_to_Z','Z_to_P','Z_to_END','P_to_Z','P_to_P','P_to_END'])],8)
    d.add_picture(str(charts['transition']),width=Inches(6.4));caption(d,'그림 7-3. 검색 상태 전이·종료 건수')
    body(d,'Treatment에서는 Z→P가 증가하고 Z→Z 및 Z→END가 감소했다. P에서 출발하는 전이는 treatment 전용 수정 대상이 아니며 경로 분기 이후 구성 차이로 건수가 달라진다. 평균 검색 수는 treatment 6.3566으로 control 6.7144보다 0.3578회 적다.')
    d.add_heading('7.6 시간분포',level=2)
    tp=data['time'];add_table(d,['자료','간격 n','동일 timestamp','>109초','p50','p90','p95','최대'],[[name,f"{v['n']:,}",pct(v['same']),pct(v['tail109']),f"{v['p50']:.1f}s",f"{v['p90']:.1f}s",f"{v['p95']:.1f}s",f"{v['max']:.1f}s"] for name,v in tp.items()],7.5)
    d.add_picture(str(charts['time']),width=Inches(6.4));caption(d,'그림 7-4. 검색 간격 중앙값과 상위 분위수(log scale)')
    body(d,'생성 데이터는 동일 timestamp와 긴 꼬리를 유지했다. 다만 control p95가 기준과 달라질 수 있고 cap 51,304초가 분포를 제한한다. 시간값은 합성 순서 제약을 검증하는 용도이며 실제 사용자의 숙고시간·이탈시간 추정에는 사용할 수 없다.')
    d.add_heading('7.7 기존 시나리오 민감도',level=2)
    sr=[];names={'search_zero_rate':'검색 0건률','immediate_recovery_transition_rate':'즉시 회복','final_recovery_rate':'최종 회복','card_h_search_rate':'Card H'}
    for _,s in scenario.iterrows():sr.append([s.scenario,names[s.metric],int(s.runs),pp(s.diff_mean),f"{pp(s.diff_min)}~{pp(s.diff_max)}"])
    add_table(d,['시나리오','지표','run 수','평균 T−C','seed 범위'],sr,7.5)
    d.add_picture(str(charts['scenario']),width=Inches(6.5));caption(d,'그림 7-5. 기존 30-seed 교정 시나리오 민감도')
    body(d,'교정 자료는 conservative 5, expected 20, optimistic 5개 run이며 각 run은 USER 2,000명이다. 시나리오 간 공통 seed가 없어 직접 paired seed 차이로 보지 않고 각 시나리오의 평균과 범위를 비교했다. expected 교정 결과와 본 10,000명 run은 서로 다른 표본 규모·seed이므로 합산하지 않았다.')

    d.add_heading('제8장. 인사이트와 실행 제안',level=1)
    section_guide(d,'결과를 제품·실험 설계로 어떻게 연결할 수 있는가?','회복 개입은 실제 로그를 갖춘 파일럿 후보지만, 기능별 효과를 분리하고 제안 노출·선택·예약 연결을 먼저 계측해야 한다.')
    add_table(d,['발견','수치 근거','가능한 설명','생성 가정 영향','적용 제안','실제 검증 방법','우선순위·확신'],[
        ['0건 이후 회복 경로 개선 신호','0건률 -6.34%p, 즉시회복 +5.12%p','Z 후 대안 제시가 반복 실패를 줄일 가능성','Z→P +0.06을 직접 설정','0건 화면 개입을 소규모 후보로','사용자 무작위 배정, Z 노출·선택·다음검색 기록','높음·합성 내 높음/현실 낮음'],
        ['최종 회복 증가','+6.71%p','즉시 및 후속 회복 누적','종료·회복 결합 모형 의존','세션 마지막 결과까지 추적','사전 관찰기간·세션 종료 규칙 고정','높음·중간'],
        ['상세진입 변화 없음','-0.27%p','회복만으로 1위 상세관심은 늘지 않을 수 있음','Card H 확률 arm 공통','회복 이후 정렬/콘텐츠 별도 실험','positive 검색 중 rank1 detail 고유검색','중간·중간'],
        ['지역 변경 행동 집중','T region_change 검색 16,626건','Z 후 지역 확장이 주요 모형 메커니즘','코드가 Z 후 강제','인접지역 UI부터 계측','제안 노출·수락·변경·회복 단계 로그','높음·개별효과 낮음'],
        ['가격/옵션별 차이 탐색','세그먼트 표의 방향 차이','초기 제약이 경로 노출을 바꿈','전용 기능 처치 없음','차기 factorial 또는 기능별 arm','사전 필터를 고정한 다중 arm','중간·낮음'],
        ['검색어 제안 판단 곤란','원본 query-change 근거 n=10','효과 분포 불안정','최종 intent·노출 없음','파일럿 계측 우선','query suggestion 노출/선택과 결과 연결','탐색·낮음'],
        ['예약 성과 판단 불가','BOOKING 0행','연결 승인 전 제외','생성 자체 금지','현재 배포 판단에서 제외','ROOM–HOTEL·search_id 브리지 검증 후 등록','보류'],
    ],6.1)
    body(d,'교수 문서의 1~4순위는 가격·옵션·지역·반복 방지였으나 실제 구현은 지역 변경과 공통 회복 uplift에 집중됐다. 합성 결과는 “0건 이후 개입 패키지를 더 검토할 가치”는 제시하지만, 가격과 옵션 중 무엇이 더 효과적인지는 말하지 못한다. 실제 서비스에서는 인접지역을 포함한 한 개입부터 노출·수락·조건변경·회복·상세진입을 계측하고, 가격/옵션/검색어를 독립 arm 또는 요인 설계로 분리해야 한다.')

    d.add_heading('제9장. 증강과 본 분석의 한계',level=1)
    section_guide(d,'1만 명이라는 숫자가 해결하지 못하는 문제는 무엇인가?','합성 표본은 Monte Carlo 변동을 줄여도 원본 소표본·모형 편향·누락 로그·인과 식별 문제를 해결하지 못한다.')
    limitations=[
        ('원본 소표본','active 41명·43세션·296검색','확률 추정 불안정','현실 일반화 제한','원본 n 병기','추가 실관측 수집'),
        ('1,000명 합성 의존','합성 기준·반복 패턴','정보량 중복','좁은 구간 오해','합성으로 표시','원본 후험·외부 검증'),
        ('0건률 16/20','요구 18, pooled +0.4551%p','seed drift','기준선 민감','HOLD 보존','독립 seed·새 관측'),
        ('5,000/arm 확대','모형 표본만 증가','Monte Carlo 오차만 감소','편향 미해소','조건부 구간','실사용자 실험'),
        ('구간 의미 차이','기준 비율 구간≠새 seed 예측','coverage 혼용 위험','PASS 오판','명칭 분리','사전 예측구간 설계'),
        ('QA PASS와 HOLD','관계 gate PASS, 통계 gate HOLD','구조는 정상·재현성 부족','신뢰 범위 혼동','두 상태 병기','별도 승인 기준'),
        ('순환적 가정','Z→P +0.06 설정','회복 증가 내재','효과 입증 불가','메커니즘 점검으로 해석','라이브 A/B'),
        ('복합 treatment','지역 변경+uplift 결합','기능 기여 미식별','우선순위 제한','패키지로만 보고','다중 arm/factorial'),
        ('사후 선택','Z 경험·회복자 조건부','분모 구성 변화','전체 효과 오해','전체 배정군 우선','ITT와 노출자 분석 사전등록'),
        ('pair·상관','pair/세션 내 반복 검색','행 독립 위반','p값 과소','pair bootstrap','cluster robust 검증'),
        ('조건부 구간','고정 모형 bootstrap','모형오류 미포함','불확실성 과소','범위 주석','다중 seed/모형 비교'),
        ('검색어 희소','원본 n=10','후험 불안정','순위 확정 불가','낮은 확신','추가 노출 표본'),
        ('시간 합성','body/tail+cap 51,304초','극단값 인공 제약','체류시간 해석 불가','순서 QA로 제한','실제 event clock 수집'),
        ('NULL5·미측정','연령·숙소유형·등급·리뷰 전부 NULL','이질성 미측정','타깃팅 제한','미실시 표시','필드 수집 승인'),
        ('노출/수락 부재','제안 상세 로그 없음','퍼널 단절','기능 수용성 불명','역추정 금지','suggestion event schema'),
        ('BOOKING 0','예약 출처/연결 미승인','성과 행 없음','매출 판단 불가','KPI 제외','ROOM–HOTEL·booking bridge'),
        ('clone 0의 한계','A/B/C 0','기계복제만 배제','현실 대표성 보장 아님','다양성 병기','외부 분포 비교'),
        ('반복 개발·seed','교정 seed는 개발/평가용','평가 적응 가능성','독립성 약화','production seed 분리','새 holdout seed/데이터'),
        ('정의 변경','Card H·회복·SG 이력','시계열 비교 단절','과거 숫자 혼합','contract 고정','실험 전 지표 사전등록'),
        ('후속검색 없음','session END만 생성','실제 이탈 직접 관측 아님','이탈률 오해','후속없음으로 표현','exit/close/time-out 로그'),
    ]
    add_table(d,['한계','확인 근거','데이터 영향','분석·의사결정 영향','이번 대응','남은 검증'],limitations,5.9)
    note(d,'BOOKING 0행은 예약전환율 0%가 아니다. control 편차 +0.4551%p를 treatment 효과에서 일괄 차감하지 않았고, treatment 차이가 이 편차보다 크다는 이유로 실제 효과를 검증했다고 판단하지 않았다. 합성 표본 증가와 실제 관측 근거 증가는 다르다.')

    d.add_heading('제10장. 결론과 후속 검증 계획',level=1)
    section_guide(d,'다음 단계에서 무엇을 실제로 검증해야 하는가?','사용자 단위 무작위 배정, 사전등록 지표, 제안 퍼널 로그, booking 연결을 갖춘 소규모 실제 A/B가 우선이다.')
    add_table(d,['범주','결론'],[
        ['이번 시뮬레이션에서 확인','고정 모형에서 10,000명 paired 데이터가 구조 오류·실제 clone 없이 생성됐고, expected treatment에서 0건·회복 지표가 설정 방향으로 움직였다. Card H는 거의 변하지 않았다.'],
        ['이번 자료로 확인하지 못함','실제 인과효과, 6개 의도군별 효과, 가격/옵션/지역/검색어 개별 기여, 제안 수락률, 이탈, 예약·매출, 실제 모집단 일반화.'],
        ['실제 서비스 다음 검증','사용자 단위 무작위 배정과 사전등록된 검색 0건률·즉시/최종 회복·Card H-SEARCH, 제안 노출/선택/조건변경 로그, booking 연결 검증.'],
    ],8)
    body(d,'권장 실험은 사용자 단위로 무작위 배정하고 동일 사용자의 검색을 한 arm에 유지한다. 대상은 최초 또는 세션 중 0건 경험 사용자로 사전 고정하되, 전체 배정군 ITT와 노출자 조건부 분석을 분리한다. 제안 노출·제안 유형·선택·실제 필터 변경·후속 결과·1위 상세진입을 search_id와 pair가 아닌 실제 experiment_assignment_id로 연결한다.')
    body(d,'주 지표는 검색 0건률 또는 Z→P 즉시 회복 중 하나를 선택하고 grain·분모·관찰기간·세션 종료·중복 이벤트 처리를 사전등록한다. 최종 회복과 Card H-SEARCH는 보조 지표로 둔다. 예약 KPI는 booking_complete–search_id–hotel_id–room_id 연결 및 중복 규칙을 승인한 뒤 포함한다.')
    body(d,'표본 설계는 합성 10,000명을 근거로 검정력이 확보됐다고 보지 않는다. 새 실제 로그의 기준률과 실무적으로 의미 있는 최소효과(MDE), 사용자 내 반복·군집효과, 탈락률을 사용해 다시 계산해야 한다. 초기 2주 계측 파일럿 후 본 실험 규모를 확정하는 순서가 안전하다.')

    d.add_heading('부록 A. 핵심 SQL·계산 절차',level=1)
    body(d,'검색 0건률: SELECT COUNT(DISTINCT CASE WHEN total_result_count=0 THEN search_id END) / COUNT(DISTINCT search_id) FROM Search. Arm 분석에서는 SessionSynthetic와 UserSynthetic를 연결해 sample_set_type으로 분리한다.')
    body(d,'즉시 회복: session_id 안에서 search_sequence 순 LEAD(total_result_count)를 계산한 뒤, 현재가 Z이고 다음 검색이 존재하는 전이를 분모, 다음 결과가 양수인 전이를 분자로 합산한다. 첫 Z 지표는 누적 Z 순번이 1인 행만 남긴다.')
    body(d,"Card H-SEARCH: 결과 존재 Search를 분모로 하고 SearchResult.result_rank=1의 hotel_id와 ActionEvent.event_type='hotel_detail_view'를 search_id·hotel_id로 연결한 뒤 DISTINCT search_id를 분자로 센다. 반복 event row는 H-EVENT로 별도 집계한다.")
    body(d,'Pair bootstrap 의사코드: for b in 1..2000: ① 5,000 pair_id를 복원추출, ② 선택된 pair의 control/treatment 전체 기여를 함께 합산, ③ 각 arm의 분자/분모를 재계산, ④ treatment−control 저장, ⑤ 2.5/97.5 분위수를 조건부 구간으로 보고한다.')
    body(d,'상태 전이 의사코드: 첫 P/Z 추출 → 현재 Z이면 first/later Z별 END·Z·P 확률 사용(실험군 P 확률에 scenario uplift 추가) → 현재 P이면 Z 경험 전/후의 P→Z 확률과 END 적용 → END 또는 31검색에서 종료 → 마지막 상태로 SG1/SG3/SG4 기록. 이는 과거 문서의 SG1~SG4와 완전히 같은 정의가 아니므로 직접 비교하지 않는다.')

    d.add_heading('부록 B. 입력·출력·버전·SHA-256',level=1)
    srcrows=[]
    for s in sources:srcrows.append([s['role'],Path(s['path']).name,s['kind'],s['sha256'],s['use']])
    add_table(d,['역할','파일명','근거 구분','SHA-256','사용 위치'],srcrows,5.7)
    size=DB10.stat().st_size;body(d,f'최종 10,000명 DB 용량은 {size:,} bytes = {size/1_000_000:.6f} MB(10진) = {size/1_048_576:.6f} MiB(2진)다. 파일 크기는 분석 전후 동일했다.')
    body(d,f"실행 seed=2434815518, scenario=expected, generation model=newref_v02+sequential_v06, config SHA-256={sha256(CONFIG)}, metric contract SHA-256={sha256(METRIC_CONTRACT)}. 분석 환경은 Python {platform.python_version()}, pandas {pd.__version__}, numpy {np.__version__}, matplotlib {matplotlib.__version__}, python-docx 사용; bootstrap seed={ANALYSIS_SEED}, 반복={BOOTSTRAP_REPS:,}다.")

    d.add_heading('부록 C. 출처 매핑과 미실시 분석',level=1)
    add_table(d,['수치·판정','근거 파일/테이블·쿼리'],[
        ['원본 41 active/43 session/296 search','v03 SQLite: event↔search, COUNT DISTINCT'],['1,000명 핵심 기준','채택 reference SQLite Search/Event/SearchResult, 공통 정의 SQL'],['30-seed 16/20·pooled·범위','260907_1418_04 handoff metric_summary 및 검증된 sidecar 30개'],['10,000명 KPI','최종 DB Search/SessionSynthetic/UserSynthetic/ActionEvent/SearchResult 직접 SQL'],['무결성·clone','전수QA JSON, raw gate, clone replay audit'],['교수 우선순위','교수판단회신 2절·표 4, Card 정정 요구 4절·표 8'],['현재 Card H 정의','공통 metric contract 및 원본/기준/10k DB 재계산'],['제한 사용','3-L authorization·제한사용기록·analysis handoff'],
    ],7)
    add_table(d,['미실시 분석','사유'],[
        ['6개 intent_segment별 효과','최종 DB에 intent_segment가 없음'],['제안 노출→선택→수락 퍼널','제안 유형·선택 필드가 없음'],['가격/옵션/검색어 개별 인과효과','복합 treatment이며 독립 arm이 없음'],['숙소유형·등급·연령·리뷰 분석','NULL5 구조적 결측'],['예약전환·매출','BOOKING 0행 및 연결 미승인'],['실제 이탈률','후속검색 없음은 이탈 직접 로그가 아님'],['10,000명 conservative/optimistic','새 합성 DB 생성 금지; 기존 2,000명/run 교정만 재사용'],['실서비스 검정력','실제 기준률·MDE·군집효과 미확정'],
    ],8)
    body(d,'최종 검수 체크: 10,000명 결과와 30-seed 교정을 분리했고, 핵심 비율은 분자/분모를 표시했다. 회복과 SG를 혼용하지 않았으며 Card H는 고유 검색 기준이다. 사후 집단을 전체 효과로 표현하지 않았고 strict HOLD와 limited-use를 함께 보존했다. 예약성과·실제 효과·소표본 신뢰도 향상을 주장하지 않았다.')
    return d


def main():
    if OUTPUT.exists():raise FileExistsError(OUTPUT)
    for path,expected in EXPECTED.items():
        if not path.exists() or sha256(path)!=expected:raise RuntimeError(f"input hash mismatch: {path}")
    before={str(p):{"sha256":sha256(p),"bytes":p.stat().st_size,"mtime_ns":p.stat().st_mtime_ns} for p in [DB10,REF,ORIGINAL,STRICT,CONFIG,GENERATOR,METRIC_CONTRACT,AUTHORITY]}
    handoff=json.loads(HANDOFF.read_text(encoding='utf-8'));qa=json.loads(QA.read_text(encoding='utf-8'));strict=json.loads(STRICT.read_text(encoding='utf-8'))
    if handoff['limited_use_status']!='LIMITED_USE_READY' or handoff['strict_calibration_status']!='HOLD' or handoff['database']['sha256']!=EXPECTED[DB10]:raise RuntimeError('handoff gate')
    if any(x['status']!='PASS' for x in qa['checks']):raise RuntimeError('QA failure')
    config=yaml.safe_load(CONFIG.read_text(encoding='utf-8'))
    orig=metric_profile(ORIGINAL);ref=metric_profile(REF);frame=session_contributions();effects=paired_bootstrap(frame);segments=segment_analysis(frame)
    time={'기준 1,000':time_profile(REF),'control':time_profile(DB10,'control'),'treatment':time_profile(DB10,'treatment')}
    filters=pd.DataFrame(filter_profiles(REF)+filter_profiles(DB10,'control')+filter_profiles(DB10,'treatment'))
    scenario=scenario_summary(strict)
    charts=make_charts(effects,segments,qa,scenario,time)
    doc_paths=[AUTHORITY,
      ROOT/'01_프로젝트기획'/'팀프로젝트'/'2026'/'08'/'일본호텔검색_팀프로젝트계획서_20260828_v03_현행본.docx',
      ROOT/'09_단계별 분석'/'2단계_1000건_증감_분석'/'09_R3-B_독립문서검수_260905_1826_01'/'호텔검색_관측형합성1000명_전체가설분석보고서_260905_1826_03.docx',
      STAGE/'호텔검색_32일차_2팀_실행체크리스트_20260905_v04_클린제출본_[TO-BE].docx',STAGE/'BI시각화_33일차_2팀프로젝트실습_해답_20260905_v04_클린제출본_[TO-BE].docx',STAGE/'BI시각화_34일차_2팀프로젝트실습_해답_20260905_v02_클린제출본_[TO-BE].docx',STAGE/'BI시각화_32-34일차_2팀실습_ASIS-TOBE_세그먼트AB재설계_20260905_v02_클린제출본.docx',STAGE/'호텔검색_1000명_10000명_세그먼트AB시뮬레이션_증강계획서_20260905_v02_클린제출본.docx',STAGE/'BI시각화_32-34일차_TO-BE_슬랙공유_브리핑_20260905_v02_클린제출본.docx']
    sources=[
      {'role':'분석 대상 DB','path':str(DB10),'kind':'10,000명 합성 결과','sha256':sha256(DB10),'use':'제7장 전체'},
      {'role':'채택 기준 DB','path':str(REF),'kind':'1,000명 합성 기준','sha256':sha256(REF),'use':'제2·6·7장'},
      {'role':'비식별 원본','path':str(ORIGINAL),'kind':'원본 관측 사실','sha256':sha256(ORIGINAL),'use':'제2·4장'},
      {'role':'분석 handoff','path':str(HANDOFF),'kind':'제한 승인','sha256':sha256(HANDOFF),'use':'입력 gate'},
      {'role':'교정 handoff','path':str(STRICT),'kind':'30-seed 교정','sha256':sha256(STRICT),'use':'제4·7장'},
      {'role':'생성기','path':str(GENERATOR),'kind':'구현 코드','sha256':sha256(GENERATOR),'use':'제3장'},
      {'role':'expected config','path':str(CONFIG),'kind':'시나리오 가정','sha256':sha256(CONFIG),'use':'제3·5장'},
      {'role':'metric contract','path':str(METRIC_CONTRACT),'kind':'정정 지표 정의','sha256':sha256(METRIC_CONTRACT),'use':'제6장'},
    ]
    for i,p in enumerate(doc_paths):sources.append({'role':'최우선 권위 문서' if i==0 else '참고 문서','path':str(p),'kind':'설계/문서 근거','sha256':sha256(p),'use':'제1~5·8장'})
    data={'handoff':handoff,'qa':qa,'strict':strict,'config':config,'orig':orig,'ref':ref,'frame':frame,'effects':effects,'segments':segments,'time':time,'filters':filters,'scenario':scenario,'charts':charts,'sources':sources}
    effects.to_json(INTERNAL/'analysis_effects.json',orient='records',force_ascii=False,indent=2);segments.to_json(INTERNAL/'analysis_segments.json',orient='records',force_ascii=False,indent=2);scenario.to_json(INTERNAL/'scenario_summary.json',orient='records',force_ascii=False,indent=2)
    document=build_document(data);document.save(OUTPUT)
    after={str(p):{"sha256":sha256(p),"bytes":p.stat().st_size,"mtime_ns":p.stat().st_mtime_ns} for p in [DB10,REF,ORIGINAL,STRICT,CONFIG,GENERATOR,METRIC_CONTRACT,AUTHORITY]}
    verification={"output":str(OUTPUT),"sha256":sha256(OUTPUT),"bytes":OUTPUT.stat().st_size,"input_unchanged":before==after,"bootstrap_seed":ANALYSIS_SEED,"bootstrap_reps":BOOTSTRAP_REPS,"effects":effects.to_dict('records'),"sections":len(document.sections),"paragraphs":len(document.paragraphs),"tables":len(document.tables),"inline_shapes":len(document.inline_shapes)}
    (INTERNAL/'build_verification.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    print(json.dumps({k:verification[k] for k in ['output','sha256','bytes','input_unchanged','paragraphs','tables','inline_shapes']},ensure_ascii=True))


if __name__=='__main__':main()
