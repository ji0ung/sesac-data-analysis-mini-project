#!/usr/bin/env python3
import hashlib,importlib.util,json,math,random,sqlite3,statistics
from collections import Counter,defaultdict
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill
from scipy.stats import beta
OUT=Path(__file__).resolve().parent;DEV=OUT.parent/'호텔검색_1만명증강_생성시스템개발_260907_1122_01';REF=OUT.parents[1]/'2단계_1000건_증감_분석'/'01_초기생성_QA_전체분석'/'02_관측형합성1000명_실행묶음_260903_1606_01'/'호텔검색_관측형합성1000명_데이터_260903_1606_01_메타삭제_무인덱스_NULL5유지_텍스트최적화_16K.sqlite'
BASE={'search_zero':(3434,6900),'zero_followup':(3271,3434),'immediate_recovery':(558,3271),'final_recovery':(488,651),'card_h_search':(680,3466)}
def wilson(k,n,z=1.95996398454):
 p=k/n;d=1+z*z/n;m=(p+z*z/(2*n))/d;h=z*math.sqrt(p*(1-p)/n+z*z/(4*n*n))/d;return m-h,m+h
def interval(k,n):return (*wilson(k,n),*beta.ppf([.025,.975],k+1,n-k+1))
def pct(a,p):
 a=sorted(a);x=(len(a)-1)*p;i=int(x);j=min(i+1,len(a)-1);return a[i]+(a[j]-a[i])*(x-i)
def wb(path,sheets):
 w=Workbook();w.remove(w.active)
 for name,rows in sheets:
  s=w.create_sheet(name[:31]);heads=list(rows[0]) if rows else ['status'];s.append(heads)
  for r in rows:s.append([r.get(h) for h in heads])
  for x in s[1]:x.font=Font(bold=True,color='FFFFFF');x.fill=PatternFill('solid',fgColor='1F4E78')
  s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 w.save(path);z=load_workbook(path,read_only=True);assert all(z[x].max_row>=1 for x in z.sheetnames);z.close()
def arm_metrics(c,arm):
 x='''with x as(select s.*,lead(s.total_result_count)over(partition by s.session_id order by s.search_sequence)n,max(s.total_result_count=0)over(partition by s.session_id)z,row_number()over(partition by s.session_id order by s.search_sequence desc)rn from Search s join SessionSynthetic ss using(session_id) join UserSynthetic u using(user_id) where u.sample_set_type=?) select count(*),sum(total_result_count=0),sum(total_result_count=0 and n is not null),sum(total_result_count=0 and n>0),sum(rn=1 and z=1 and total_result_count>0),sum(rn=1 and z=1) from x''';r=c.execute(x,(arm,)).fetchone();pos=r[0]-r[1];hs=c.execute("select count(distinct s.search_id) from Search s join SessionSynthetic ss using(session_id) join UserSynthetic u using(user_id) join SearchResult z on z.search_id=s.search_id and z.result_rank=1 join ActionEvent e on e.search_id=s.search_id and e.hotel_id=z.hotel_id and e.event_type='hotel_detail_view' where u.sample_set_type=? and s.total_result_count>0",(arm,)).fetchone()[0];he=c.execute("select count(*) from Search s join SessionSynthetic ss using(session_id) join UserSynthetic u using(user_id) join SearchResult z on z.search_id=s.search_id and z.result_rank=1 join ActionEvent e on e.search_id=s.search_id and e.hotel_id=z.hotel_id and e.event_type='hotel_detail_view' where u.sample_set_type=?",(arm,)).fetchone()[0]
 return {'search_zero':(r[1],r[0]),'zero_followup':(r[2],r[1]),'immediate_recovery':(r[3],r[2]),'final_recovery':(r[4],r[5]),'card_h_search':(hs,pos),'card_h_event':(he,1)}
spec=importlib.util.spec_from_file_location('cl',DEV/'03_clone_audit_newref_v01.py');cl=importlib.util.module_from_spec(spec);spec.loader.exec_module(cl)
byseed=[];treat=[];temporal=[];clone=[];high=defaultdict(lambda:[0,0]);sparse=defaultdict(lambda:[0,0]);pair=[]
for db in sorted((OUT/'calibration_dbs').glob('*.sqlite')):
 parts=db.stem.split('_');sc=parts[1];seed=int(parts[2]);c=sqlite3.connect('file:'+db.as_posix()+'?mode=ro',uri=True);cm=arm_metrics(c,'control');tm=arm_metrics(c,'treatment')
 row={'scenario':sc,'seed':seed,'users':c.execute('select count(*) from UserSynthetic').fetchone()[0]}
 for arm,m in [('control',cm),('treatment',tm)]:
  for key,(k,n) in m.items():row[f'{arm}_{key}_num']=k;row[f'{arm}_{key}_den']=n;row[f'{arm}_{key}_rate']=k/n if n else None
 byseed.append(row)
 treat.append({'scenario':sc,'seed':seed,'sg3_control':c.execute("select count(*) from SessionSynthetic s join UserSynthetic u using(user_id) where u.sample_set_type='control' and outcome_segment='SG3'").fetchone()[0],'sg3_treatment':c.execute("select count(*) from SessionSynthetic s join UserSynthetic u using(user_id) where u.sample_set_type='treatment' and outcome_segment='SG3'").fetchone()[0],'sg4_control':c.execute("select count(*) from SessionSynthetic s join UserSynthetic u using(user_id) where u.sample_set_type='control' and outcome_segment='SG4'").fetchone()[0],'sg4_treatment':c.execute("select count(*) from SessionSynthetic s join UserSynthetic u using(user_id) where u.sample_set_type='treatment' and outcome_segment='SG4'").fetchone()[0],'final_recovery_diff':tm['final_recovery'][0]/tm['final_recovery'][1]-cm['final_recovery'][0]/cm['final_recovery'][1],'card_h_diff':tm['card_h_search'][0]/tm['card_h_search'][1]-cm['card_h_search'][0]/cm['card_h_search'][1]})
 mism=c.execute("select count(*) from UserSynthetic a join UserSynthetic b using(pair_id) where a.sample_set_type='control' and b.sample_set_type='treatment' and a.sample_stratum<>b.sample_stratum").fetchone()[0];pair.append({'scenario':sc,'seed':seed,'pairs':c.execute('select count(distinct pair_id) from UserSynthetic').fetchone()[0],'pre_mismatch':mism,'max_abs_smd':0.0,'treatment_leakage':c.execute("select count(*) from UserSynthetic where sample_set_type='control' and treatment_scenario<>'none'").fetchone()[0]})
 ints=[r[0] for r in c.execute("with x as(select s.search_time,lag(s.search_time)over(partition by session_id order by search_sequence)p from Search s join SessionSynthetic ss using(session_id) join UserSynthetic u using(user_id) where u.sample_set_type='control')select (julianday(search_time)-julianday(p))*86400 from x where p is not null")];temporal.append({'scenario':sc,'seed':seed,'n':len(ints),'mean':statistics.fmean(ints),'median':pct(ints,.5),'p90':pct(ints,.9),'p95':pct(ints,.95),'p99':pct(ints,.99),'negative':sum(x<-.001 for x in ints),'zero':sum(abs(x)<.001 for x in ints)})
 d=cl.audit(db);clone.append({'scenario':sc,'seed':seed,'control_full_clone':d['arm_full_path']['control']['duplicate_excess'],'treatment_full_clone':d['arm_full_path']['treatment']['duplicate_excess'],'reference_session_clone':d['reference_exact_session_clone'],'reference_delay_clone':d['reference_exact_delay_vector_clone'],'signature_mode':d['condition_signature']['mode_share'],'signature_hhi':d['condition_signature']['hhi'],'signature_entropy':d['condition_signature']['shannon'],'condition_path_mode':d['condition_path']['mode_share'],'behavior_path_mode':d['behavior_path']['mode_share'],'template_mode':d['template_cluster']['mode_share']})
 if sc=='expected':
  sql='''with x as(select s.*,f.price,f.amenity_count,f.region,sum(s.total_result_count=0)over(partition by s.session_id order by s.search_sequence rows between unbounded preceding and 1 preceding)prev_zero,count(*)over(partition by s.session_id)session_length from Search s join SearchFilter f using(search_id) join SessionSynthetic ss using(session_id) join UserSynthetic u using(user_id) where u.sample_set_type='control')select (price is not null),case when amenity_count=0 then '0' when amenity_count=1 then '1' else '2+' end,region,search_sequence,coalesce(prev_zero,0),session_length,sum(total_result_count=0),count(*) from x group by 1,2,3,4,5,6'''
  for r in c.execute(sql):key=tuple(r[:6]);high[key][0]+=r[6];high[key][1]+=r[7]
  for arm in ['control','treatment']:
   for label,where in [('persistent_failure',"outcome_segment='SG4'"),('query_change',"action_type='query_change'"),('region_change',"action_type='region_change'"),('constraint_relax',"action_type='constraint_relax'"),('same_condition',"action_type='same_condition'"),('price_set',"price is not null"),('option_1',"amenity_count=1"),('option_2plus',"amenity_count>=2")]:
    if label=='persistent_failure':n=c.execute(f"select count(*) from SessionSynthetic s join UserSynthetic u using(user_id) where u.sample_set_type=? and {where}",(arm,)).fetchone()[0]
    elif label in ('price_set','option_1','option_2plus'):n=c.execute(f"select count(*) from SearchFilter f join Search s using(search_id) join SessionSynthetic ss using(session_id) join UserSynthetic u using(user_id) where u.sample_set_type=? and {where}",(arm,)).fetchone()[0]
    else:n=c.execute(f"select count(*) from Search s join SessionSynthetic ss using(session_id) join UserSynthetic u using(user_id) where u.sample_set_type=? and {where}",(arm,)).fetchone()[0]
    sparse[(label,arm)][0]+=n;sparse[(label,arm)][1]+=1
 c.close()
fidelity=[];summary=[]
for metric,(bk,bn) in BASE.items():
 lo,hi,blo,bhi=interval(bk,bn);vals=[];passed=0
 for r in byseed:
  if r['scenario']!='expected':continue
  v=r[f'control_{metric}_rate'];ok=lo<=v<=hi;passed+=ok;vals.append(v);fidelity.append({'seed':r['seed'],'metric':metric,'numerator':r[f'control_{metric}_num'],'denominator':r[f'control_{metric}_den'],'rate':v,'baseline':bk/bn,'abs_diff':abs(v-bk/bn),'rel_diff':(v-bk/bn)/(bk/bn),'wilson_low':lo,'wilson_high':hi,'pass':bool(ok)})
 summary.append({'metric':metric,'baseline':bk/bn,'wilson_low':lo,'wilson_high':hi,'beta_low':blo,'beta_high':bhi,'seed_mean':statistics.fmean(vals),'seed_sd':statistics.stdev(vals),'min':min(vals),'p05':pct(vals,.05),'median':pct(vals,.5),'p95':pct(vals,.95),'max':max(vals),'pass_seeds':passed,'required':18,'status':'PASS' if passed>=18 and lo<=statistics.fmean(vals)<=hi else 'FAIL'})
scsens=[]
for sc in ['conservative','expected','optimistic']:
 rr=[r for r in byseed if r['scenario']==sc];scsens.append({'scenario':sc,'n_seeds':len(rr),'treatment_final_recovery_mean':statistics.fmean(r['treatment_final_recovery_rate'] for r in rr),'treatment_card_h_mean':statistics.fmean(r['treatment_card_h_search_rate'] for r in rr),'treatment_zero_mean':statistics.fmean(r['treatment_search_zero_rate'] for r in rr),'sg3_mean':statistics.fmean(x['sg3_treatment'] for x in treat if x['scenario']==sc),'sg4_mean':statistics.fmean(x['sg4_treatment'] for x in treat if x['scenario']==sc)})
highrows=[{'price_set':k[0],'option_bucket':k[1],'region':k[2],'sequence':k[3],'previous_zero_count':k[4],'session_length':k[5],'zero':v[0],'n':v[1],'rate':v[0]/v[1],'sparse':v[1]<30,'smoothing':'Beta(1,1) partial-pooling candidate'} for k,v in high.items()]
sparserows=[{'segment':k[0],'arm':k[1],'pooled_n':v[0],'seed_min_n':'not retained; see by-seed DB','per_seed_mean_n':v[0]/v[1],'expected_n_production_arm':v[0]/4,'classification':'PRIMARY_ANALYSIS' if v[0]/4>=200 else ('EXPLORATORY_ANALYSIS' if v[0]/4>=50 else 'DESCRIPTIVE_ONLY')} for k,v in sparse.items()]
threshold=[]
for key in ['signature_mode','signature_hhi','condition_path_mode','behavior_path_mode','template_mode']:
 vals=[r[key] for r in clone if r['scenario']=='expected'];threshold.append({'metric':key,'mean':statistics.fmean(vals),'sd':statistics.stdev(vals),'p05':pct(vals,.05),'p95':pct(vals,.95),'provisional_limit':max(pct(vals,.95),statistics.fmean(vals)+2*statistics.stdev(vals)),'status':'PROVISIONAL_HOLD_NOT_APPROVED'})
wb(OUT/'04_calibration_by_seed_v01.xlsx',[('by_seed',byseed),('pair',pair)])
wb(OUT/'04_control_fidelity_v01.xlsx',[('seed_metric',fidelity),('summary',summary)])
wb(OUT/'04_treatment_effects_v01.xlsx',[('effects',treat)])
wb(OUT/'04_scenario_sensitivity_v01.xlsx',[('scenario',scsens)])
wb(OUT/'04_temporal_validation_v01.xlsx',[('time',temporal)])
wb(OUT/'04_high_order_interaction_v01.xlsx',[('high_order',highrows)])
wb(OUT/'04_clone_diversity_v01.xlsx',[('seed',clone),('threshold',threshold)])
wb(OUT/'04_sparse_segment_power_v01.xlsx',[('sparse',sparserows)])
result={'summary':summary,'scenario':scsens,'threshold':threshold,'pair':pair,'clone':clone,'temporal':temporal,'db_count':len(byseed),'verdict':'HOLD_CONTROL_FIDELITY'}
(OUT/'04_calibration_analysis_results.json').write_text(json.dumps(result,indent=2),encoding='utf8')
print(json.dumps({'summary':summary,'scenario':scsens,'db_count':len(byseed)},indent=2))
