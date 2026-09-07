#!/usr/bin/env python3
"""PROMPT 2.8 authoritative-reference audit. Opens SQLite read-only."""
import hashlib,json,math,sqlite3,statistics,sys
from collections import Counter,defaultdict
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill

DB,OUT=map(Path,sys.argv[1:3]); TS='260907_1110'; TABLES=['user','hotel','room','search','search_filter','search_result','event','booking']
PK={'user':'user_id','hotel':'hotel_id','room':'room_id','search':'search_id','search_filter':'search_filter_id','search_result':'search_result_id','event':'event_id','booking':'booking_id'}
FKS=[('room','hotel_id','hotel','hotel_id'),('search_filter','search_id','search','search_id'),('search_result','search_id','search','search_id'),('search_result','hotel_id','hotel','hotel_id'),('search_result','room_id','room','room_id'),('event','hotel_id','hotel','hotel_id'),('event','search_filter_id','search_filter','search_filter_id'),('event','search_id','search','search_id'),('event','user_id','user','user_id'),('booking','user_id','user','user_id'),('booking','hotel_id','hotel','hotel_id'),('booking','room_id','room','room_id')]
NULL5=[('event','review_completed_at'),('event','review_text'),('search_filter','property_type'),('search_filter','property_grade'),('user','age_group')]
def sha(p):
 h=hashlib.sha256()
 with p.open('rb') as f:
  for b in iter(lambda:f.read(1<<20),b''):h.update(b)
 return h.hexdigest()
def dt(x): return datetime.strptime(x.replace(' KST',''),'%Y-%m-%d %H:%M:%S')
def qtile(a,p):
 if not a:return None
 a=sorted(a);z=(len(a)-1)*p;l=int(z);h=min(l+1,len(a)-1);return a[l]+(a[h]-a[l])*(z-l)
def stats(a):
 return {'n':len(a),'min':min(a) if a else None,'p25':qtile(a,.25),'median':qtile(a,.5),'mean':statistics.fmean(a) if a else None,'p75':qtile(a,.75),'p90':qtile(a,.9),'p95':qtile(a,.95),'max':max(a) if a else None}
def diversity(vals):
 c=Counter(vals);n=sum(c.values());ps=[v/n for v in c.values()] if n else []
 return {'n':n,'unique':len(c),'mode_count':max(c.values()) if c else 0,'mode_share':max(ps) if ps else 0,'hhi':sum(x*x for x in ps),'shannon_entropy':-sum(x*math.log(x) for x in ps),'duplicate_excess':sum(v-1 for v in c.values())}
def workbook(path,sheets):
 w=Workbook();w.remove(w.active)
 for name,rows in sheets:
  s=w.create_sheet(name[:31]); heads=list(rows[0]) if rows else ['status'];s.append(heads)
  for row in rows:s.append([row.get(h) for h in heads])
  for x in s[1]:x.font=Font(bold=True,color='FFFFFF');x.fill=PatternFill('solid',fgColor='1F4E78')
  s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 w.save(path)

c=sqlite3.connect('file:'+DB.as_posix()+'?mode=ro',uri=True);c.row_factory=sqlite3.Row;c.execute('pragma query_only=on')
objects={r[0] for r in c.execute("select name from sqlite_master where type='table'")}; integrity=c.execute('pragma integrity_check').fetchone()[0];quick=c.execute('pragma quick_check').fetchone()[0]
rows={t:c.execute(f'select count(*) from "{t}"').fetchone()[0] for t in TABLES}
qa=[]
def add(name,value,expected=0):qa.append({'check':name,'value':value,'expected':expected,'status':'PASS' if value==expected else 'FAIL'})
add('missing_business_tables',len(set(TABLES)-objects))
for t,k in PK.items():
 add(f'{t}.{k}.null',c.execute(f'select count(*) from "{t}" where "{k}" is null').fetchone()[0]);add(f'{t}.{k}.duplicate_excess',c.execute(f'select count(*)-count(distinct "{k}") from "{t}"').fetchone()[0])
for a,ac,b,bc in FKS:add(f'{a}.{ac}->{b}.{bc}.orphan',c.execute(f'select count(*) from "{a}" a left join "{b}" b on a."{ac}"=b."{bc}" where a."{ac}" is not null and b."{bc}" is null').fetchone()[0])
add('search_filter_not_1to1',c.execute('select count(*) from (select search_id,count(*) n from search_filter group by search_id having n<>1) x').fetchone()[0]+c.execute('select count(*) from search s left join search_filter f using(search_id) where f.search_id is null').fetchone()[0])
add('result_count_mismatch',c.execute('select count(*) from search s left join(select search_id,count(*) n from search_result group by search_id)r using(search_id) where s.total_result_count<>coalesce(r.n,0)').fetchone()[0])
add('duplicate_hotel_within_search',c.execute('select count(*) from(select search_id,hotel_id,count(*) n from search_result group by search_id,hotel_id having n>1)').fetchone()[0])
add('event_search_session_mismatch',c.execute('select count(*) from event e join search s using(search_id) where e.search_id is not null and e.session_id<>s.session_id').fetchone()[0])
add('session_multiple_users',c.execute('select count(*) from(select session_id,count(distinct user_id)n from event where user_id is not null group by session_id having n<>1)').fetchone()[0])
add('search_session_without_user',c.execute('select count(*) from(select distinct session_id from search except select distinct session_id from event where user_id is not null)').fetchone()[0])
add('search_sequence_duplicate',c.execute("select count(*) from(select session_id,cast(substr(search_id,length(search_id)-2)as int) seq,count(*) n from search group by session_id,seq having n>1)").fetchone()[0])
add('search_sequence_time_inversion',c.execute("with x as(select session_id,search_time,lag(search_time)over(partition by session_id order by cast(substr(search_id,length(search_id)-2)as int)) p from search)select count(*) from x where search_time<p").fetchone()[0])
add('event_before_search',c.execute('select count(*) from event e join search s using(search_id) where e.search_id is not null and e.event_at<s.search_time').fetchone()[0])
add('session_end_before_event',c.execute('select count(*) from event where session_end_time is not null and session_end_time<event_at').fetchone()[0])
add('booking_rows',rows['booking'])
for t,col in NULL5:add(f'{t}.{col}.not_null',c.execute(f'select count(*) from "{t}" where "{col}" is not null').fetchone()[0])

total=rows['search'];zero=c.execute('select count(*) from search where total_result_count=0').fetchone()[0];pos=total-zero
follow=c.execute("with x as(select total_result_count,lead(search_id)over(partition by session_id order by search_time,search_id)n from search)select sum(total_result_count=0 and n is not null),sum(total_result_count=0)from x").fetchone()
imm=c.execute("with x as(select total_result_count,lead(total_result_count)over(partition by session_id order by search_time,search_id)n from search)select sum(total_result_count=0 and n>0),sum(total_result_count=0 and n is not null)from x").fetchone()
final=c.execute("with x as(select session_id,total_result_count,row_number()over(partition by session_id order by search_time desc,search_id desc)rn,max(total_result_count=0)over(partition by session_id)z from search)select sum(rn=1 and z=1 and total_result_count>0),sum(rn=1 and z=1)from x").fetchone()
hd=c.execute("select count(distinct s.search_id) from search s join search_result r on r.search_id=s.search_id and r.result_rank=1 join event e on e.search_id=s.search_id and e.hotel_id=r.hotel_id and e.event_type='hotel_detail_view' where s.total_result_count>0").fetchone()[0]
he=c.execute("select count(*) from search s join search_result r on r.search_id=s.search_id and r.result_rank=1 join event e on e.search_id=s.search_id and e.hotel_id=r.hotel_id and e.event_type='hotel_detail_view' where s.total_result_count>0").fetchone()[0]
experienced=final[1];persistent=experienced-final[0]
kpis=[
 {'metric':'search_zero','numerator':zero,'denominator':total,'rate':zero/total}, {'metric':'positive_search','numerator':pos,'denominator':total,'rate':pos/total},
 {'metric':'zero_followup','numerator':follow[0],'denominator':follow[1],'rate':follow[0]/follow[1]}, {'metric':'immediate_recovery','numerator':imm[0],'denominator':imm[1],'rate':imm[0]/imm[1]},
 {'metric':'final_recovery','numerator':final[0],'denominator':final[1],'rate':final[0]/final[1]}, {'metric':'persistent_failure_among_zero_sessions','numerator':persistent,'denominator':experienced,'rate':persistent/experienced},
 {'metric':'session_final_failure_all','numerator':c.execute('with x as(select total_result_count,row_number()over(partition by session_id order by search_time desc,search_id desc)rn from search)select sum(rn=1 and total_result_count=0)from x').fetchone()[0],'denominator':1000,'rate':persistent/1000},
 {'metric':'card_h_search','numerator':hd,'denominator':pos,'rate':hd/pos},{'metric':'card_h_event','numerator':he,'denominator':1,'rate':float(he)},{'metric':'card_h_repeat_mean','numerator':he,'denominator':hd,'rate':he/hd}]
session_len=[dict(r) for r in c.execute('select count(*) searches,count(*) sessions from search group by session_id').fetchall()] # replaced below
session_len=[{'searches':r[0],'sessions':r[1]} for r in c.execute('select searches,count(*) from(select session_id,count(*) searches from search group by session_id)group by searches order by searches')]
hazard=[{'sequence':r[0],'zero':r[1],'total':r[2],'rate':r[1]/r[2]} for r in c.execute("select cast(substr(search_id,length(search_id)-2)as int),sum(total_result_count=0),count(*) from search group by 1 order by 1")]

searches=defaultdict(list)
for r in c.execute('select * from search order by session_id,search_time,search_id'):searches[r['session_id']].append(r)
sf={r['search_id']:r for r in c.execute('select * from search_filter')}; results=defaultdict(list);events=defaultdict(list)
for r in c.execute('select * from search_result order by search_id,result_rank,hotel_id,room_id'):results[r['search_id']].append(r)
for r in c.execute('select * from event where search_id is not null order by search_id,event_at,event_id'):events[r['search_id']].append(r)
inter=[];zero_inter=[];detail_inter=[];click_detail=[];first_zero_research=[];condition_sig=[];condition_paths=[];behavior_paths=[];full_paths=[];template_clusters=[]
for sid,ss in searches.items():
 condition=[];behavior=[];full=[];start=dt(ss[0]['search_time']); hotelmap={};roommap={}
 firstzero=next((i for i,s in enumerate(ss) if s['total_result_count']==0),None)
 if firstzero is not None:first_zero_research.append(len(ss)-firstzero-1)
 for i,s in enumerate(ss):
  st=dt(s['search_time']);
  if i: inter.append((st-dt(ss[i-1]['search_time'])).total_seconds())
  if i and ss[i-1]['total_result_count']==0:zero_inter.append((st-dt(ss[i-1]['search_time'])).total_seconds())
  f=sf[s['search_id']]; sig=(s['query_text'],s['checkin_date'],s['checkout_date'],s['destination'],s['guest_count'],s['sort_option'],f['property_type'],f['property_grade'],f['user_rating_min'],f['price'],f['amenity_count'],f['region']);condition_sig.append(sig);condition.append(sig)
  evtypes=tuple(e['event_type'] for e in events[s['search_id']]);behavior.append(('zero' if s['total_result_count']==0 else 'positive',evtypes))
  rr=[]
  for x in results[s['search_id']]:
   hotelmap.setdefault(x['hotel_id'],f'H{len(hotelmap)+1}');roommap.setdefault(x['room_id'],f'R{len(roommap)+1}');rr.append((x['result_rank'],hotelmap[x['hotel_id']],roommap[x['room_id']]))
  ee=[]
  for e in events[s['search_id']]:ee.append((e['event_type'],(dt(e['event_at'])-start).total_seconds(),hotelmap.get(e['hotel_id'],'NONE')))
  full.append(((st-start).total_seconds(),sig,s['total_result_count'],tuple(rr),tuple(ee)))
  if s['total_result_count']>0:
   top=next((x['hotel_id'] for x in results[s['search_id']] if x['result_rank']==1),None); ds=[dt(e['event_at']) for e in events[s['search_id']] if e['event_type']=='hotel_detail_view' and e['hotel_id']==top]
   if ds:detail_inter.append((min(ds)-st).total_seconds())
  clicks=defaultdict(list);details=defaultdict(list)
  for e in events[s['search_id']]:
   if e['event_type']=='hotel_click':clicks[e['hotel_id']].append(dt(e['event_at']))
   if e['event_type']=='hotel_detail_view':details[e['hotel_id']].append(dt(e['event_at']))
  for h,aa in clicks.items():
   bb=details.get(h,[])
   for a in aa:
    b=next((x for x in bb if x>=a),None)
    if b is not None:click_detail.append((b-a).total_seconds())
 condition_paths.append(tuple(condition));behavior_paths.append(tuple(behavior));full_paths.append(tuple(full));template_clusters.append((len(ss),tuple(x[0] for x in behavior)))

time_rows=[]
for name,a in [('search_interarrival_seconds',inter),('zero_to_research_seconds',zero_inter),('positive_to_top_detail_seconds',detail_inter),('click_to_detail_seconds',click_detail)]:
 d=stats(a);d.update({'metric':name,'negative':sum(x<0 for x in a),'zero':sum(x==0 for x in a),'same_timestamp_rate':sum(x==0 for x in a)/len(a) if a else None,'mode_share':max(Counter(a).values())/len(a) if a else None});time_rows.append(d)
positive_inter=[x for x in inter if x>0];logs=[math.log(x) for x in positive_inter]
time_rows.append({'metric':'positive_interarrival_lognormal_reference','n':len(logs),'mean':statistics.fmean(logs),'median':None,'min':min(logs),'p25':qtile(logs,.25),'p75':qtile(logs,.75),'p90':qtile(logs,.9),'p95':qtile(logs,.95),'max':max(logs),'negative':0,'zero':0,'same_timestamp_rate':0,'mode_share':None,'log_mu':statistics.fmean(logs),'log_sigma_sample':statistics.stdev(logs)})
time_summary={'search_time_min':c.execute('select min(search_time) from search').fetchone()[0],'search_time_max':c.execute('select max(search_time) from search').fetchone()[0],'intervals':time_rows,'sequence_time_inversions':next(x['value'] for x in qa if x['check']=='search_sequence_time_inversion')}
diversity_rows=[]
for name,vals in [('condition_signature',condition_sig),('condition_path',condition_paths),('behavior_path',behavior_paths),('template_cluster_derived_session_length_and_outcome_sequence',template_clusters),('full_session_path',full_paths)]:
 d=diversity(vals);d['definition']=name;diversity_rows.append(d)
full=next(x for x in diversity_rows if x['definition']=='full_session_path')
clone={'full_path_duplicate_groups':sum(v>1 for v in Counter(full_paths).values()),'duplicate_related_sessions':sum(v for v in Counter(full_paths).values() if v>1),'full_path_clone_excess':full['duplicate_excess'],'classification':'REFERENCE_TEMPLATE_REPETITION_NOT_FOR_REPRODUCTION'}

result={'db':{'path':str(DB.resolve()),'bytes':DB.stat().st_size,'sha256':sha(DB),'integrity':integrity,'quick':quick,'page_size':c.execute('pragma page_size').fetchone()[0]},'rows':rows,'qa':qa,'kpis':kpis,'session_length':session_len,'hazard':hazard,'first_zero_research_count_distribution':dict(sorted(Counter(first_zero_research).items())),'time':time_summary,'diversity':diversity_rows,'clone':clone,'null5':NULL5}
workbook(OUT/f'호텔검색_1만명증강_신규기준선프로파일_{TS}_01.xlsx',[('rows',[{'table':k,'rows':v} for k,v in rows.items()]),('logical_integrity',qa),('kpi',kpis),('session_length',session_len),('sequence_hazard',hazard),('zero_research_dist',[{'following_searches':k,'sessions':v} for k,v in sorted(Counter(first_zero_research).items())])])
workbook(OUT/f'호텔검색_1만명증강_시간간격감사_{TS}_01.xlsx',[('time_summary',time_rows),('session_bounds',[{'session_id':sid,'first_search':ss[0]['search_time'],'last_search':ss[-1]['search_time']} for sid,ss in searches.items()])])
workbook(OUT/f'호텔검색_1만명증강_경로반복분류_{TS}_01.xlsx',[('diversity',diversity_rows),('classification',[{'category':'REFERENCE_OBSERVED_PATTERN','action':'preserve conditional distributions, hazards, session length and event intensity'}, {'category':'REFERENCE_TEMPLATE_REPETITION','action':'do not copy rows, sessions, full paths or source_session_id payloads'}]),('clone',[clone])])
(OUT/'audit_results.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
c.close();print(json.dumps({'qa_failures':[x for x in qa if x['status']=='FAIL'],'kpis':kpis,'time':time_summary,'diversity':diversity_rows,'clone':clone},ensure_ascii=False))
